"""Synthetic capability oracles, independent of production names and coordinates."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.table import Table


def workbook(sheets, formats=None):
    book = Workbook()
    book.remove(book.active)
    for number, (name, matrix) in enumerate(sheets.items(), 1):
        sheet = book.create_sheet(name)
        for row in matrix:
            sheet.append(row)
        sheet.add_table(Table(displayName=f"Data{number}", ref=sheet.dimensions))
    for name, address, style in formats or []:
        book[name][address].number_format = style
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def capability_cases():
    join = workbook({
        "Orders": [["Record", "Customer ID", "Amount"], ["ORD-1", "C1", 10],
                   ["ORD-2", "C2", 20], ["ORD-3", "C1", 30]],
        "Customers": [["Customer ID", "Region"], ["C1", "North"], ["C2", "South"]],
    })
    ambiguous = workbook({
        "North": [["Record", "Amount"], ["Entry Alpha", 10], ["Entry Beta", 20]],
        "South": [["Record", "Amount"], ["Entry Gamma", 70], ["Entry Delta", 80]],
    })
    formulas = workbook({"Calculations": [
        ["Record", "Quantity", "Unit Price", "Amount"],
        ["Entry Alpha", 2, 10, "=B2*C2"], ["Entry Beta", 3, 7, "=B3*C3"],
        ["Entry Gamma", 1, 9, "=B4*C4"],
    ]})
    external = workbook({"Unavailable": [["Record", "Amount"], ["Entry Alpha", "=IFERROR('[1]Remote'!B2,0)"]]})
    cross = workbook({"Inputs": [["Record", "Count"], ["Alpha", 2], ["Beta", 3], ["Gamma", 4]],
                      "Derived": [["Record", "Amount"], ["Total", "=SUM(Inputs!B2:B4)*2"]]})
    percent = workbook({"Derived": [["Record", "Ratio"], ["Quarter", "=1/4"]]}, [("Derived", "B2", "0.0%")])
    duplicate = workbook({
        "Orders": [["Record", "Customer ID", "Amount"], ["ORD-1", "C1", 10], ["ORD-2", "C2", 20]],
        "Customers": [["Customer ID", "Region"], ["C1", "North"], ["C1", "South"]],
    })
    return [
        ("join_group", join, "Orders와 Customers를 Customer ID로 연결해서 Region별 Amount 합계", "answered",
         [r"North[^\n]*40", r"South[^\n]*20"], "source_checked"),
        ("join_lookup", join, "Orders와 Customers를 Customer ID로 연결해서 ORD-2의 Region과 Amount 알려줘", "answered",
         [r"South", r"20"], "source_checked"),
        ("join_duplicate", duplicate, "Orders와 Customers를 Customer ID로 연결해서 Region별 Amount 합계", "insufficient_evidence",
         [r"중복"], "unverified"),
        ("ambiguous_scope", ambiguous, "Amount 합계는 얼마야?", "insufficient_evidence",
         [r"North", r"South", r"다시 질문"], "unverified"),
        ("resolved_scope", ambiguous, "South 시트의 Amount 합계를 알려줘", "answered",
         [r"150"], "request_checked"),
        ("formula_aggregate", formulas, "Calculations 시트의 Amount 합계를 알려줘", "limited",
         [r"50"], "source_checked"),
        ("formula_external", external, "'Unavailable'!B2 셀 값은 얼마야?", "insufficient_evidence",
         [r"외부", r"저장"], "unverified"),
        ("formula_cross_sheet", cross, "'Derived'!B2 셀 값은 얼마야?", "limited",
         [r"18"], "source_checked"),
        ("formula_percent", percent, "'Derived'!B2 셀 값은 얼마야?", "limited",
         [r"25(?:\.0)?%"], "source_checked"),
    ]
