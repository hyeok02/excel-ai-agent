"""Restore cited tail cells without claiming that a sparse index is complete."""
from openpyxl.utils import get_column_letter

from app.agent.query.references import MAX_CITATION_AREA, _box, _contains, normalize_reference
from app.agent.query.source_context_matches import context_matches
from app.agent.query.source_lookup_reader import read_source_range
from app.agent.query.source_lookup_references import SourceRange, bind_source_reference
from app.agent.query.workbook_source import source_for
from app.agent.query.workbook_source_models import SourceReadError

MAX_CONTEXT_READS = 6
MAX_CONTEXT_READ_CELLS = 768


def refresh_context_cells(context, index, cells, tables):
    if getattr(index, "source_provider", None) is None:
        return cells
    source = source_for(index)
    if source is None:
        return {}
    reader = ContextReadback(source, dict(cells))
    items = list(reversed(context.turns)) + list(context.insights)
    for item in items:
        reader.ensure(item.evidence)
        keys = {key for reference in item.evidence for key in context_matches(reference, reader.cells, index)}
        matched = [reader.cells[key] for key in sorted(keys)]
        related = []
        for table in tables:
            bounds = _box(table.id)
            selected = [cell for cell in matched if bounds and (box := _box(cell.reference)) and _contains(bounds, box)]
            if not selected:
                continue
            columns = {box.min_column for cell in selected if (box := _box(cell.reference))}
            ordered = sorted(table.columns, key=lambda column: not (
                (box := _box(column.header_reference)) and box.min_column in columns))
            related.extend(column.header_reference for column in ordered[:6])
            related.extend(reference for reference, _ in table.context[:2])
            related.extend(f"{table.sheet_name}!{column.id}{box.min_row}" for cell in selected[:12]
                           if (box := _box(cell.reference)) for column in table.columns[:12])
        reader.ensure(related)
    return reader.cells


class ContextReadback:
    def __init__(self, source, cells):
        self.source, self.cells = source, cells
        self.reads = self.area = 0

    def ensure(self, references):
        requested = []
        for reference in dict.fromkeys(references):
            try:
                item = bind_source_reference(reference, self.source)
            except SourceReadError:
                continue
            if item.area > min(MAX_CITATION_AREA, MAX_CONTEXT_READ_CELLS):
                continue
            if all(normalize_reference(f"{item.sheet}!{address}") in self.cells for address in item.addresses()):
                continue
            requested.append(item)
        for item in _coalesce(requested):
            if self.reads >= MAX_CONTEXT_READS or self.area + item.area > MAX_CONTEXT_READ_CELLS:
                break
            self.reads += 1
            self.area += item.area
            try:
                found = read_source_range(self.source, item, MAX_CONTEXT_READ_CELLS)
            except SourceReadError:
                continue
            self.cells.update((normalize_reference(cell.reference), cell) for cell in found
                              if cell.value is not None or cell.formula)


def _coalesce(requested):
    groups = {}
    for item in requested:
        left, top, right, bottom = item.bounds
        key = item.sheet, top, bottom
        existing = groups.get(key)
        if existing:
            left, right = min(left, existing.bounds[0]), max(right, existing.bounds[2])
        merged = SourceRange(item.sheet, f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
                             (left, top, right, bottom))
        if merged.area <= MAX_CONTEXT_READ_CELLS:
            groups[key] = merged
        else:
            yield item
    yield from groups.values()
