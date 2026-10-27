"""Read-back proofs for numeric row-label/column-label matrix intersections.

Three nearby cells are not a matrix. Require a contiguous numeric rectangle
with at least two labelled rows and columns, drawn from the current full index.
Only current tool evidence receives bindings; duplicate label pairs are retained
as alternatives so the claim checker cannot arbitrarily pick one.
"""
from collections import defaultdict
from math import isfinite

from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from app.agent.query.references import normalize_reference


def source_matrix_bindings(summary, index, evidence):
    if index is None:
        return {}
    original = {normalize_reference(cell.reference): cell for row in index.rows for cell in row.cells}
    candidates = {}
    for item in evidence:
        reference = normalize_reference(f"{item.sheet_name}!{item.reference}")
        cell = original.get(reference)
        if (cell and type(cell.value) is type(item.value) and cell.value == item.value
                and cell.formula == item.formula):
            candidates[reference] = cell
    # Search all indexed sheets for competing label pairs, not just the rows
    # the model happened to retrieve. Retrieval is not proof of uniqueness.
    sheets = defaultdict(dict)
    for row in index.rows:
        sheets[row.sheet_name][row.row_number] = {
            coordinate_to_tuple(cell.address)[1]: cell for cell in row.cells
        }
    merged = {sheet.name: {span for region in sheet.regions for span in region.merged_ranges}
              for sheet in summary.sheets}
    layouts = []
    for sheet, rows in sheets.items():
        if index.truncated_for({sheet}):
            continue
        layouts.extend(_layouts(sheet, rows, merged.get(sheet, ())))
    alternatives = defaultdict(set)
    by_cell = defaultdict(list)
    for layout in layouts:
        for record in layout:
            key = (_label(record["row_value"]), _label(record["column_value"]))
            alternatives[key].add((record["sheet"], record["value_reference"]))
            by_cell[record["value_reference"]].append(record)
    bindings = {}
    for reference, records in by_cell.items():
        if reference not in candidates:
            continue
        axes = {(item["row_reference"], item["column_reference"]) for item in records}
        if len(axes) != 1:
            continue
        record = records[0]
        if not all(ref in candidates for ref in (record["row_reference"], record["column_reference"])):
            continue
        key = (_label(record["row_value"]), _label(record["column_value"]))
        bindings[reference] = {**record,
            "sheet_names": [sheet.name for sheet in summary.sheets],
            "unavailable_sheets": sorted(index.truncated_sheet_names) if index.truncated else [],
            "alternatives": [
            {"sheet": sheet, "reference": ref} for sheet, ref in sorted(alternatives[key])
        ]}
    return bindings


def _layouts(sheet, rows, merged):
    for top, headers in rows.items():
        for left, header in sorted(headers.items()):
            if left < 2 or not _text(header) or not _text(headers.get(left + 1)):
                continue
            # A blank corner distinguishes independent column labels from an
            # ordinary textual record (for example "Has leases? / No / note")
            # sitting above numeric records. Do not promote that data record
            # to a header merely because nearby numbers form a rectangle.
            corner = headers.get(left - 1)
            if corner and (corner.formula or corner.value not in (None, "")):
                continue
            first = rows.get(top + 1, {})
            if not _text(first.get(left - 1)) or not all(_number(first.get(col)) for col in (left, left + 1)):
                continue
            right = left + 1
            while _text(headers.get(right + 1)):
                right += 1
            records, bottom = [], top
            while True:
                row = rows.get(bottom + 1, {})
                if not _text(row.get(left - 1)) or not all(_number(row.get(col)) for col in (left, left + 1)):
                    break
                row_right = left + 1
                while row_right < right and _number(row.get(row_right + 1)):
                    row_right += 1
                right = min(right, row_right)
                bottom += 1
                records.append(row)
            if len(records) < 2 or _merged_matrix(merged, left - 1, top, right, bottom):
                continue
            matrix_range = f"{get_column_letter(left - 1)}{top}:{get_column_letter(right)}{bottom}"
            yield [{
                "sheet": sheet, "matrix_range": matrix_range,
                "value_reference": normalize_reference(row[col].reference),
                "row_reference": normalize_reference(row[left - 1].reference),
                "column_reference": normalize_reference(headers[col].reference),
                "value": row[col].value, "row_value": row[left - 1].value,
                "column_value": headers[col].value,
            } for row in records for col in range(left, right + 1)]


def _merged_matrix(spans, left, top, right, bottom):
    for span in spans:
        a, b, c, d = range_boundaries(span)
        if a <= right and left <= c and b <= bottom and top <= d and (a != c or b != d):
            return True
    return False


def _number(cell):
    return cell is not None and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and isfinite(cell.value)


def _text(cell):
    return cell is not None and not cell.formula and isinstance(cell.value, str) and bool(cell.value.strip())


def _label(value):
    return " ".join(str(value).casefold().split())
