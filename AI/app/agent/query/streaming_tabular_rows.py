"""Re-read source headers and rows without treating absent samples as blanks."""
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from app.agent.query.ambiguity_labels import source_mentioned
from app.agent.query.ambiguity_qualifiers import explicit_column_ids
from app.agent.query.tabular_catalog_models import QueryRow
from app.agent.query.tabular_filters import matches
from app.agent.query.tabular_request_scope import header_request_text
from app.agent.query.tabular_metadata_headers import MAX_METADATA_ROWS, detail_metadata_row, series_metadata_rows


def source_scope(provider, index, table, needed, tables, question):
    if (not index.source_digest or provider.source_digest != index.source_digest
            or table.sheet_name not in provider.authorized_sheets
            or table.source_kind not in {"native", "schema", "region", "report"}
            or not table.complete and table.source_kind == "report"):
        return None
    sheet = provider.sheet(table.sheet_name)
    if sheet is None:
        return None
    left, top, right, bottom = range_boundaries(table.reference)
    if top in sheet.hidden_rows or right > sheet.columns or bottom > sheet.rows:
        return None
    hidden = needed & (sheet.hidden_columns | table.hidden_columns)
    explicit = explicit_column_ids(question)
    field_text = header_request_text(question, {candidate.sheet_name for candidate in tables})
    if any(column.id in hidden and column.id not in explicit and not source_mentioned(column.header, field_text)
           for column in table.columns):
        return None
    if not _external_headers_match(provider, sheet, table, top):
        return None
    for span in sheet.merged_ranges:
        a, b, c, d = range_boundaries(span)
        if a <= right and left <= c and b <= bottom and top < d and (a != c or b != d):
            return None
    for other in tables:
        if other.id == table.id or other.sheet_name != table.sheet_name:
            continue
        a, b, c, d = range_boundaries(other.reference)
        if a <= right and left <= c and b <= bottom and top <= d:
            return None
    return sheet


def header_matches(row, table):
    cells = {cell.address: cell for cell in row.cells}
    for column in table.columns:
        if coordinate_to_tuple(column.header_reference.rpartition("!")[2])[0] < row.row_number:
            continue  # A vertical merged anchor was read in source_scope.
        source = cells.get(column.header_reference.rpartition("!")[2])
        if source is None or not isinstance(source.value, str) or source.value.strip() != column.header:
            return False
    return True


def _external_headers_match(provider, sheet, table, top):
    expected = {}
    merged = [range_boundaries(span) for span in sheet.merged_ranges]
    for column in table.columns:
        address = column.header_reference.rpartition("!")[2]
        row, number = coordinate_to_tuple(address)
        if row != top:
            if row > top or not any(left == right == number and start == row and end >= top
                                    for left, start, right, end in merged):
                return False
            expected[address] = column.header
        for reference, label in column.header_path:
            anchor = reference.rpartition("!")[2]
            ancestor_row, ancestor_column = coordinate_to_tuple(anchor)
            if not any(left == ancestor_column and start == ancestor_row
                       and left <= number <= right and end < row for left, start, right, end in merged):
                return False
            expected[anchor] = label
        for reference, identifier in column.source_identifiers:
            address = reference.rpartition("!")[2]
            row, number = coordinate_to_tuple(address)
            if (reference.rpartition("!")[0] != table.sheet_name or row != top + 1
                    or get_column_letter(number) != column.id or row not in table.metadata_rows):
                return False
            expected[address] = identifier
    if not expected:
        return True
    coordinates = [coordinate_to_tuple(address) for address in expected]
    first_row = min(row for row, _ in coordinates)
    last_row = max(row for row, _ in coordinates)
    first_column = min(column for _, column in coordinates)
    last_column = max(column for _, column in coordinates)
    reference = f"{get_column_letter(first_column)}{first_row}:{get_column_letter(last_column)}{last_row}"
    actual = {}
    scan = provider.scan_rows(table.sheet_name, reference)
    with scan as rows:
        for row in rows:
            if row.sheet_name != table.sheet_name:
                return False
            actual.update((cell.address, cell.value) for cell in row.cells)
    return scan.complete and all(isinstance(actual.get(address), str)
                                 and actual[address].strip() == label for address, label in expected.items())


def metadata_matches(row, table):
    cells = {cell.address: {"value": cell.value, "formula": cell.formula, "cached_value": cell.value}
             for cell in row.cells}
    cells.update({column.header_reference.rpartition("!")[2]: {"value": column.header}
                  for column in table.columns})
    left, top, right, _ = range_boundaries(table.reference)
    if row.row_number == top + 1:
        return series_metadata_rows(table.reference, cells) == frozenset({row.row_number})
    return (top + 1 in table.metadata_rows and top + 2 <= row.row_number <= top + MAX_METADATA_ROWS
            and row.row_number in table.metadata_rows
            and detail_metadata_row(cells, left, row.row_number, right))


def source_row(row, table):
    by_column = {get_column_letter(coordinate_to_tuple(cell.address)[1]): cell for cell in row.cells}
    values = {column.id: by_column[column.id].value if column.id in by_column else None for column in table.columns}
    formulas = {column: cell.formula for column, cell in by_column.items() if column in values and cell.formula}
    formats = {column: cell.number_format or "General" for column, cell in by_column.items()
               if column in values and cell.value not in (None, "")}
    return QueryRow(row.row_number, values, formulas, formats)


def row_eligible(row, plan, table):
    if getattr(plan, "filter_expression", None) is not None:
        from app.agent.query.boolean_filters import evaluate_filters
        return evaluate_filters(row, plan.filters, {column.id: column for column in table.columns}, plan.filter_expression)
    headers = {column.id: column.header for column in table.columns}
    for condition in plan.filters:
        result = matches(row.values[condition.column_id], condition,
                         row.number_formats.get(condition.column_id, "General"), headers[condition.column_id])
        if result is None:
            return None
        if not result:
            return False
    return True


