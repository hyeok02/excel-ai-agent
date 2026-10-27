"""Build typed source tables from validated rectangular cell collections."""
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from app.agent.query.tabular_catalog_models import QueryColumn, QueryRow, QueryTable
from app.agent.query.tabular_catalog_sources import cell_value
from app.agent.query.tabular_metadata_headers import series_identifiers, series_metadata_rows


def build_table(sheet, reference, cells, complete, merged, hidden):
    left, top, right, bottom = range_boundaries(reference)
    if bottom <= top or right <= left:
        return None
    headers, header_references = [], {}
    observed_rows = sorted({coordinate_to_tuple(address)[0] for address in cells
                            if top < coordinate_to_tuple(address)[0] <= bottom})
    for number in range(left, right + 1):
        column = get_column_letter(number)
        raw = cells.get(f"{column}{top}", {})
        anchor = f"{column}{top}"
        if not raw:
            for span in merged:
                a, b, c, d = range_boundaries(span)
                if a == c == number and b < top <= d:
                    anchor = f"{column}{b}"
                    raw = cells.get(anchor, {})
                    break
        value = cell_value(raw)
        if value is None:
            continue
        # Editable openpyxl loads may materialize a link target in an empty cell.
        # A hyperlink without observations below it is not a table attribute.
        if raw.get("hyperlink") == value and not any(
            cell_value(cells.get(f"{column}{row}", {})) not in (None, "")
            or cells.get(f"{column}{row}", {}).get("formula")
            for row in observed_rows
        ):
            continue
        if not isinstance(value, str) or not value.strip():
            return None
        headers.append((column, value.strip()))
        header_references[column] = f"{sheet}!{anchor}"
    if len(headers) < 2:
        return None
    if any(_data_merge(span, top + 1, bottom, headers) for span in merged):
        return None
    metadata = series_metadata_rows(reference, cells)
    rows = []
    for number in observed_rows:
        if number in metadata:
            continue
        values = {column: cell_value(cells.get(f"{column}{number}", {}))
                  for column, _ in headers}
        formulas = {column: raw["formula"] for column, _ in headers
                    if (raw := cells.get(f"{column}{number}", {})).get("formula")}
        formats = {column: str(cells.get(f"{column}{number}", {}).get("number_format") or "General")
                   for column, value in values.items() if value not in (None, "")}
        if not formulas and not any(value not in (None, "") for value in values.values()):
            continue
        rows.append(QueryRow(number, values, formulas, formats))
    if not rows:
        return None
    columns = tuple(QueryColumn(
        column, header, header_references[column], _samples(rows, column),
        tuple(dict.fromkeys(row.number_formats[column] for row in rows if column in row.number_formats)),
        source_identifiers=series_identifiers(sheet, reference, cells, column, metadata),
    ) for column, header in headers)
    return QueryTable(f"{sheet}!{reference}", sheet, reference, columns, tuple(rows), complete, hidden,
                      metadata_rows=metadata)


def _data_merge(span, top, bottom, headers):
    a, b, c, d = range_boundaries(span)
    columns = {get_column_letter(number) for number in range(a, c + 1)}
    selected = [column for column, _ in headers if column in columns]
    return b <= bottom and d >= top and bool(selected) and (d > b or len(selected) > 1)


def _samples(rows, column):
    values, seen = [], set()
    for row in rows:
        value = row.values[column]
        key = (type(value), str(value))
        if key not in seen:
            seen.add(key)
            values.append(value[:160] if isinstance(value, str) else value)
        if len(values) == 6:
            break
    return tuple(values)
