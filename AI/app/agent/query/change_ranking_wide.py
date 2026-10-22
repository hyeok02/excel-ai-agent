"""Bind grouped measure columns using original hierarchical headers."""
from dataclasses import replace
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from app.agent.query.comparison_scope import GROUP_ALIASES
from app.agent.query.tabular_catalog_models import QueryColumn, QueryRow
from app.agent.query.tabular_catalog_sources import _overlaps
from app.agent.query.trend_source_readback import table_rows
from app.agent.query.trend_source_vertical import vertical_candidate
from app.agent.query.trend_source_models import TrendCandidate


def wide_ranking_candidates(question, summary, index, tables, group):
    result = []
    aliases = GROUP_ALIASES.get(group, set())
    if not aliases:
        return result
    sheets = {sheet.name: sheet for sheet in summary.sheets}
    for table in tables:
        sheet = sheets[table.sheet_name]
        left, _, right, _ = range_boundaries(table.reference)
        schemas = [schema for schema in sheet.column_schemas
                   if _overlaps(schema.source_range, table.reference)
                   and left <= coordinate_to_tuple(schema.column.upper() + "1")[1] <= right
                   and len(schema.header_path) > 1
                   and any(str(part).casefold().strip() in aliases for part in schema.header_path[:-1])]
        if not schemas:
            continue
        try:
            normalized, ancestors = _normalized_table(table, tables, schemas, sheet, index, question)
        except ValueError as error:
            result.append(TrendCandidate(table.id, (), error=str(error)))
            continue
        selected = {schema.column.upper() for schema in schemas}
        discovery = question + " " + " ".join(column.header for column in normalized.columns
                                               if column.id in selected)
        for metric in normalized.columns:
            if metric.id not in selected:
                continue
            candidate = vertical_candidate(normalized, metric, [normalized, *tables], index, discovery)
            refs = ancestors.get(metric.id, ())
            candidate = replace(candidate, context=tuple(dict.fromkeys((*candidate.context, *refs))))
            result.append(candidate)
    return result


def _normalized_table(table, tables, schemas, sheet, index, question):
    rows, error = table_rows(table, tables, index, question)
    if error:
        raise ValueError(error)
    left, top, right, bottom = range_boundaries(table.reference)
    cells = {cell.address: cell for row in rows for cell in row.cells}
    all_cells = {cell.address: cell for row in index.rows if row.sheet_name == table.sheet_name
                 for cell in row.cells}
    all_cells.update(cells)
    chosen, ancestors = {}, {}
    for schema in schemas:
        leaf = str(schema.header_path[-1]).strip().casefold()
        matches = [cell for cell in cells.values()
                   if get_column_letter(coordinate_to_tuple(cell.address)[1]) == schema.column.upper()
                   and top <= coordinate_to_tuple(cell.address)[0] <= min(top + 12, bottom - 1)
                   and str(cell.value).strip().casefold() == leaf]
        if len(matches) != 1:
            raise ValueError("ambiguous_group_header")
        chosen[schema.column.upper()] = matches[0]
        ancestors[schema.column.upper()] = _ancestor_cells(schema, matches[0], sheet, all_cells)
    header_rows = {coordinate_to_tuple(cell.address)[0] for cell in chosen.values()}
    if len(header_rows) != 1:
        raise ValueError("ambiguous_group_header")
    header = next(iter(header_rows))
    columns = []
    for number in range(left, right + 1):
        key = get_column_letter(number)
        cell = cells.get(f"{key}{header}")
        if cell is None or not isinstance(cell.value, str) or not cell.value.strip():
            continue
        values = [cells.get(f"{key}{row.row_number}") for row in rows if row.row_number > header]
        columns.append(QueryColumn(key, cell.value.strip(), cell.reference,
                                   tuple(item.value for item in values if item is not None)[:6],
                                   tuple(dict.fromkeys(item.number_format for item in values if item))))
    records = []
    for row in rows:
        if row.row_number <= header:
            continue
        found = {get_column_letter(coordinate_to_tuple(cell.address)[1]): cell for cell in row.cells}
        records.append(QueryRow(row.row_number,
            {column.id: found[column.id].value if column.id in found else None for column in columns},
            {key: cell.formula for key, cell in found.items() if cell.formula},
            {key: cell.number_format for key, cell in found.items()}))
    reference = f"{get_column_letter(left)}{header}:{get_column_letter(right)}{bottom}"
    normalized = replace(table, id=f"{table.sheet_name}!{reference}", reference=reference,
                         columns=tuple(columns), rows=tuple(records), complete=True)
    return normalized, ancestors


def _ancestor_cells(schema, leaf, sheet, cells):
    row, column = coordinate_to_tuple(leaf.address)
    spans = [range_boundaries(span) for region in sheet.regions for span in region.merged_ranges]
    result = []
    for part in schema.header_path[:-1]:
        matches = []
        for cell in cells.values():
            r, c = coordinate_to_tuple(cell.address)
            if not 0 < row - r <= 12 or str(cell.value).strip().casefold() != str(part).strip().casefold():
                continue
            merged = any(a == c and b == r and a <= column <= z and b <= r <= d
                         for a, b, z, d in spans)
            if c == column or merged:
                matches.append(cell)
        if not matches:
            raise ValueError("source_mismatch")
        result.append(max(matches, key=lambda cell: coordinate_to_tuple(cell.address)[0]))
    return tuple(result)
