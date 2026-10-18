"""Read a rectangle one row at a time; only exhaustion proves full coverage."""
from itertools import zip_longest

from openpyxl.utils import get_column_letter, range_boundaries

from app.agent.query.index import IndexedRow, _indexed_cells
from app.agent.query.workbook_source_models import SourceReadError
from app.services.workbook_loading import close_workbook, load_workbook_pair


class SourceScan:
    def __init__(self, source, sheet, reference):
        self.source, self.sheet, self.reference = source, sheet, reference
        self.complete = False
        self.scanned_rows = self.scanned_cells = 0
        self._books = None
        self._started = None
        self._iterated = False
        self._finished = False
        self._iterators = ()
        try:
            left, top, right, bottom = range_boundaries(reference)
            if (None in (left, top, right, bottom) or not 1 <= left <= right <= 16384
                    or not 1 <= top <= bottom <= 1048576):
                raise ValueError()
        except (TypeError, ValueError):
            raise SourceReadError("invalid_range") from None
        self.bounds = left, top, right, bottom
        area = (right - left + 1) * (bottom - top + 1)
        if area > source.limits.max_cells or bottom - top + 1 > source.limits.max_rows:
            raise SourceReadError("resource_limit")

    def __enter__(self):
        if self._started is not None:
            raise SourceReadError("already_consumed")
        self._started = self.source._budget.begin()
        try:
            self._books = load_workbook_pair(self.source.content, read_only=True)
            if any(self.sheet.name not in book.sheetnames for book in self._books):
                raise SourceReadError("missing_sheet")
            self.source._budget.check(self._started)
        except Exception as error:
            self.__exit__(None, None, None)
            if isinstance(error, SourceReadError):
                raise
            raise SourceReadError("unreadable_source") from error
        return self

    def __iter__(self):
        if self._books is None or self._iterated:
            raise SourceReadError("inactive_scan")
        self._iterated = True
        left, top, right, bottom = self.bounds
        width = right - left + 1
        iterators = [book[self.sheet.name].iter_rows(
            min_row=top, max_row=bottom, min_col=left, max_col=right,
        ) for book in self._books]
        self._iterators = iterators
        try:
            for number, pair in enumerate(zip_longest(*iterators), top):
                self.source._budget.check(self._started, width)
                first, second = pair
                if first is None or second is None or len(first) != width or len(second) != width:
                    raise SourceReadError("incomplete_range")
                for offset, cells in enumerate(zip(first, second), left):
                    expected = f"{get_column_letter(offset)}{number}"
                    for cell in cells:
                        address = getattr(cell, "coordinate", None)
                        if address is not None and address != expected or address is None and cell.value is not None:
                            raise SourceReadError("inconsistent_cell_coordinate")
                self.scanned_rows += 1
                self.scanned_cells += width
                yield IndexedRow(self.sheet.name, number, tuple(_indexed_cells(self.sheet.name, first, second)))
            if self.scanned_rows != bottom - top + 1:
                raise SourceReadError("incomplete_range")
            self.source._budget.check(self._started)
            self.complete = True
        except SourceReadError:
            raise
        except Exception as error:
            raise SourceReadError("unreadable_source") from error
        finally:
            for iterator in iterators:
                iterator.close()

    def __exit__(self, *args):
        for iterator in self._iterators:
            iterator.close()
        if self._books is not None:
            for book in self._books:
                close_workbook(book)
            self._books = None
        if self._started is not None and not self._finished:
            self.source._budget.finish(self._started)
            self._finished = True
        return False
