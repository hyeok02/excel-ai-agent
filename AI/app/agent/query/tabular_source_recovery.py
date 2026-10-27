"""Recover adjacent source header/body blocks without a workbook-name rule."""
import re

from openpyxl.utils import get_column_letter, range_boundaries

from app.agent.query.tabular_metadata_headers import series_metadata_rows


def _value(cells, column, row):
    raw = cells.get(f"{get_column_letter(column)}{row}", {})
    return raw.get("cached_value") if raw.get("formula") else raw.get("value")


def numeric_body_header(reference, cells):
    """A region starting with numbers can have its text header one row above.

Require two numeric columns and consecutive numeric body rows. This does not
promote arbitrary textual records or guess a header across a blank separator.
"""
    left, top, right, bottom = range_boundaries(reference)
    if top < 2 or bottom <= top:
        return reference
    numeric = lambda value: isinstance(value, (int, float)) and not isinstance(value, bool)
    columns = [number for number in range(left, right + 1)
               if _value(cells, number, top) not in (None, "")]
    if (len(columns) < 2 or columns != list(range(columns[0], columns[-1] + 1))
            or not all(numeric(_value(cells, number, top)) for number in columns)
            or not all(numeric(_value(cells, number, top + 1)) for number in columns)
            or not all(isinstance(_value(cells, number, top - 1), str)
                       and _value(cells, number, top - 1).strip() for number in columns)):
        return reference
    return f"{get_column_letter(columns[0])}{top - 1}:{get_column_letter(columns[-1])}{bottom}"


def series_block_ranges(regions, cells):
    """Join only adjacent same-width regions with a labeled series prefix.

The publisher's descriptive/code/unit/date bands may be split into separate
regions by formatting. Their original coordinates and bottom bound remain the
source of truth; missing samples never become evidence of complete coverage.
"""
    bounds = sorted({range_boundaries(f"{region.start_cell}:{region.end_cell}")
                     for region, _ in regions}, key=lambda item: item[1])
    for left, top, right, bottom in bounds:
        if right <= left:
            continue
        prefix = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{top + 8}"
        metadata = series_metadata_rows(prefix, cells)
        if len(metadata) < 2:
            continue
        end = bottom
        for a, b, c, d in bounds:
            if a == left and c == right and b == end + 1:
                end = d
        first = max(metadata) + 1
        if end <= first:
            continue
        if not all(_period(_value(cells, left, row))
                   and any(isinstance(_value(cells, number, row), (int, float))
                           and not isinstance(_value(cells, number, row), bool)
                           for number in range(left + 1, right + 1))
                   for row in (first, first + 1)):
            continue
        yield f"{get_column_letter(left)}{top}:{get_column_letter(right)}{end}"


def _period(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return 1900 <= value <= 2100 and value == int(value)
    return isinstance(value, str) and bool(re.fullmatch(
        r"(?:19|20)\d{2}(?:\s*(?:Q[1-4]|[A-Za-z]{3,9})|[-/.]\d{1,2}(?:[-/.]\d{1,2})?)?",
        value.strip(), re.I))
