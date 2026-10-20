"""Find the immediately preceding original date without validating stale values."""
from openpyxl.utils import coordinate_to_tuple, get_column_letter

from app.agent.query.metric_observation_cells import _human_label, _column
from app.agent.query.metric_observation_periods import source_period
from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.tabular_catalog_sources import table_sources
from app.agent.query.trend_source_readback import table_rows, rectangle_rows, column_cells
from app.agent.query.trend_source_vertical import _identity_records
from app.agent.query.workbook_source import source_for


def previous_start(candidates, summary, index, question):
    """Date-only readback chooses a boundary, never supplies answer operands.

    The normal extractor subsequently rereads and validates the entire chosen
    interval, including the immediately preceding value even when it is bad.
    """
    if not candidates or any(not candidate.observations or candidate.error for candidate in candidates):
        return None
    starts = []
    for candidate in candidates:
        first = min(candidate.observations, key=lambda point: point.period)
        value_row, _ = coordinate_to_tuple(first.value.address)
        date_row, _ = coordinate_to_tuple(first.period_cell.address)
        dates = (_vertical_dates(candidate, summary, index, question) if value_row == date_row else
                 _horizontal_dates(candidate, summary, index))
        if dates is None:
            return None
        starts.append(max((date for date in dates if date < first.period), default=first.period))
    return min(starts)


def _vertical_dates(candidate, summary, index, question):
    tables = build_query_tables(summary)
    table = next((table for table in tables if table.id == candidate.table_id), None)
    if table is None:
        return None
    rows, error = table_rows(table, tables, index, question)
    if error:
        return None
    point = candidate.observations[0]
    header, metric_number = coordinate_to_tuple(point.label.address)
    date_column = get_column_letter(coordinate_to_tuple(point.period_cell.address)[1])
    hidden = set(table.excluded_rows)
    if provider := source_for(index):
        hidden.update(provider.sheet(table.sheet_name).hidden_rows)
    records = [column_cells(row) for row in rows if row.row_number > header and row.row_number not in hidden]
    records, _, error = _identity_records(table, records, get_column_letter(metric_number), date_column, question)
    if error:
        return None
    dates = [source_period(row[date_column].value) if date_column in row else None for row in records]
    return dates if all(date is not None for date in dates) else None


def _horizontal_dates(candidate, summary, index):
    sheet_name, _, reference = candidate.table_id.rpartition("!")
    sheet = next((sheet for sheet in summary.sheets if sheet.name == sheet_name), None)
    if sheet is None:
        return None
    source = next((source for source in table_sources(sheet) if source[0] == reference), None)
    if source is None:
        return None
    rows, error = rectangle_rows(sheet_name, reference, source[1], source[2], index)
    if error:
        return None
    point = candidate.observations[0]
    axis_row = coordinate_to_tuple(point.period_cell.address)[0]
    value_row = coordinate_to_tuple(point.value.address)[0]
    axis = next((row for row in rows if row.row_number == axis_row), None)
    row = next((row for row in rows if row.row_number == value_row), None)
    if axis is None or row is None:
        return None
    dates = {_column(cell): source_period(cell.value) for cell in axis.cells}
    left = _column(point.label)
    right = min((_column(cell) for cell in row.cells if _human_label(cell)
                 and _column(cell) > left and not dates.get(_column(cell))), default=16385)
    occupied = {_column(cell) for cell in row.cells if left < _column(cell) < right
                and (cell.value not in (None, "") or cell.formula)}
    occupied.update(column for column, date in dates.items() if left < column < right and date is not None)
    result = [dates.get(column) for column in occupied]
    return result if all(date is not None for date in result) else None
