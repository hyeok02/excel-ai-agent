"""Cell selection and complete-row segmentation for recovered source tables."""
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries


def source_range(reference, regions):
    bounds = range_boundaries(reference)
    cells, complete, merged, hidden = {}, True, [], set()
    for region, values in regions:
        selected = {key: value for key, value in values.items() if _inside(key, bounds)}
        if selected:
            cells.update(selected)
            complete = complete and region_complete(region)
            merged.extend(region.merged_ranges)
            hidden.update(row for row in region.hidden_rows if bounds[1] <= row <= bounds[3])
            for span in region.merged_ranges:
                left, top, right, bottom = range_boundaries(span)
                if left == right and top < bounds[1] <= bottom and bounds[0] <= left <= bounds[2]:
                    address = f"{get_column_letter(left)}{top}"
                    if address in values:
                        cells[address] = values[address]
    return cells, complete, merged, frozenset(hidden)


def region_complete(region):
    query_complete = getattr(region, "query_complete", None)
    return region.analysis_complete if query_complete is None else query_complete


def row_segments(reference, cells, hidden_rows):
    left, top, right, bottom = range_boundaries(reference)
    rows = sorted({coordinate_to_tuple(address)[0] for address in cells})
    groups = []
    for row in rows:
        gap = range(groups[-1][-1] + 1, row) if groups else ()
        if not groups or any(number not in hidden_rows for number in gap):
            groups.append([])
        groups[-1].append(row)
    for group in groups:
        if len(group) < 2:
            continue
        start, end = group[0], group[-1]
        if group is groups[0] and all(row in hidden_rows for row in range(top, start)):
            start = top
        if group is groups[-1] and all(row in hidden_rows for row in range(end + 1, bottom + 1)):
            end = bottom
        span = f"{get_column_letter(left)}{start}:{get_column_letter(right)}{end}"
        yield span, {key: value for key, value in cells.items()
                     if start <= coordinate_to_tuple(key)[0] <= end}


def _inside(address, bounds):
    row, column = coordinate_to_tuple(address)
    left, top, right, bottom = bounds
    return left <= column <= right and top <= row <= bottom
