"""Expose source headers and bounded samples for schema-grounded query planning."""
import re
from dataclasses import replace

from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries

from app.agent.query.tabular_catalog_models import QueryColumn, QueryRow, QueryTable
from app.agent.query.tabular_catalog_sources import cell_value, source_cells, table_context, table_sources
from app.agent.query.column_context import column_context

MAX_CATALOG_TABLES = 18
MAX_CATALOG_COLUMNS = 80
MAX_CATALOG_CHARACTERS = 30_000


def build_query_tables(summary) -> list[QueryTable]:
    tables = []
    for sheet in summary.sheets:
        sheet_cells = {address: cell for region in sheet.regions
                       for address, cell in source_cells(region).items()}
        hidden_columns = frozenset(column for region in sheet.regions for column in region.hidden_columns)
        parents = tuple(span for region in sheet.regions for span in region.merged_ranges)
        for reference, cells, complete, merged, hidden, kind in table_sources(sheet):
            table = _build_table(sheet.name, reference, cells, complete, merged, hidden)
            if table is not None:
                columns = tuple(replace(column, context=column_context(sheet.name, column, sheet_cells, parents, reference))
                                for column in table.columns)
                tables.append(replace(table, columns=columns, source_kind=kind, context=table_context(sheet.name, reference, sheet_cells),
                                      hidden_columns=hidden_columns & {c.id for c in table.columns}))
    return tables


def query_catalog(summary, question: str) -> list[dict]:
    tables = build_query_tables(summary)
    tokens = set(re.findall(r"[\w]+", question.casefold()))
    tables.sort(key=lambda table: -_relevance(table, tokens))
    catalog, size = [], 0
    for table in tables[:MAX_CATALOG_TABLES]:
        entry = {
            "table_id": table.id, "sheet_name": table.sheet_name,
            "range": table.reference, "complete": table.complete,
            "row_count": len(table.rows) if table.complete else None,
            "sampled_row_count": len(table.rows),
            "excluded_row_count": len(table.excluded_rows),
            "context": [{"reference": reference, "text": value}
                        for reference, value in table.context],
            "columns": [{
                "column_id": column.id, "header": column.header,
                "header_reference": column.header_reference,
                "parent_context": list(column.context),
                "hidden": column.id in table.hidden_columns,
                "samples": list(column.sample_values),
            } for column in table.columns[:MAX_CATALOG_COLUMNS]],
        }
        entry_size = len(str(entry))
        if size + entry_size > MAX_CATALOG_CHARACTERS:
            continue
        catalog.append(entry)
        size += entry_size
    return catalog


def _build_table(sheet, reference, cells, complete, merged, hidden):
    left, top, right, bottom = range_boundaries(reference)
    if bottom <= top or right <= left:
        return None
    headers, header_references = [], {}
    observed_rows = sorted({coordinate_to_tuple(address)[0] for address in cells
                            if top < coordinate_to_tuple(address)[0] <= bottom})
    for number in range(left, right + 1):
        column = get_column_letter(number)
        raw = cells.get(f"{column}{top}", {})
        anchor = f"{column}{top}"
        if not raw:
            for span in merged:
                a, b, c, d = range_boundaries(span)
                if a == c == number and b < top <= d:
                    anchor = f"{column}{b}"
                    raw = cells.get(anchor, {})
                    break
        value = cell_value(raw)
        if value is None:
            continue
        # openpyxl's editable loader materializes hyperlink targets into empty
        # cells; its read-only loader correctly has no underlying cell value.
        # A link with no observations below it is not a source table attribute.
        if raw.get("hyperlink") == value and not any(
            cell_value(cells.get(f"{column}{row}", {})) not in (None, "")
            or cells.get(f"{column}{row}", {}).get("formula")
            for row in observed_rows
        ):
            continue
        if not isinstance(value, str) or not value.strip():
            return None
        headers.append((column, value.strip()))
        header_references[column] = f"{sheet}!{anchor}"
    if len(headers) < 2:
        return None
    if any(_data_merge(span, top + 1, bottom, headers) for span in merged):
        return None
    rows = []
    for number in observed_rows:
        values = {column: cell_value(cells.get(f"{column}{number}", {}))
                  for column, _ in headers}
        formulas = {column: raw["formula"] for column, _ in headers
                    if (raw := cells.get(f"{column}{number}", {})).get("formula")}
        formats = {column: str(cells.get(f"{column}{number}", {}).get("number_format") or "General")
                   for column, value in values.items() if value not in (None, "")}
        if not formulas and not any(value not in (None, "") for value in values.values()):
            continue
        rows.append(QueryRow(number, values, formulas, formats))
    if not rows:
        return None
    columns = tuple(QueryColumn(
        column, header, header_references[column], _samples(rows, column),
        tuple(dict.fromkeys(row.number_formats[column] for row in rows if column in row.number_formats))
    ) for column, header in headers)
    return QueryTable(f"{sheet}!{reference}", sheet, reference, columns, tuple(rows), complete, hidden)


def _data_merge(span, top, bottom, headers):
    a, b, c, d = range_boundaries(span)
    columns = {get_column_letter(number) for number in range(a, c + 1)}
    selected = [column for column, _ in headers if column in columns]
    return b <= bottom and d >= top and bool(selected) and (d > b or len(selected) > 1)


def _samples(rows, column):
    values, seen = [], set()
    for row in rows:
        value = row.values[column]
        key = (type(value), str(value))
        if key not in seen:
            seen.add(key)
            values.append(value[:160] if isinstance(value, str) else value)
        if len(values) == 6:
            break
    return tuple(values)


def _relevance(table, tokens):
    source = " ".join([table.sheet_name] + [column.header for column in table.columns])
    source += " " + " ".join(value for _, value in table.context)
    source += " " + " ".join(str(value) for column in table.columns
                             for value in column.sample_values if value is not None)
    relevance = sum(len(token) for token in tokens if len(token) > 1 and token in source.casefold())
    # Helper dictionaries remain addressable, but should not displace visible
    # report populations when an otherwise equally relevant table exists.
    return relevance - (10 if len(table.hidden_columns) == len(table.columns) else 0)
