"""Recover rectangular sources without relying on semantic role classifications."""
from collections import defaultdict
from dataclasses import asdict, is_dataclass
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries
from app.agent.query.tabular_report_sources import leaf_header_row, report_ranges
from app.agent.query.tabular_metadata_headers import recover_series_header_range
from app.agent.query.tabular_source_recovery import numeric_body_header, series_block_ranges
from app.agent.query.tabular_source_ranges import (
    region_complete as _complete, row_segments as _row_segments, source_range as _source_range,
)


def source_cells(region) -> dict[str, dict]:
    if getattr(region, "query_complete", None) is not None:
        rows = region.query_rows
    elif region.analysis_complete or region.analysis_rows:
        rows = region.analysis_rows
    else:
        rows = region.preview_rows
    return {
        str(cell["address"]): cell
        for row in rows
        for item in row
        if (cell := asdict(item) if is_dataclass(item) else item).get("address")
    }


def cell_value(cell: dict):
    return cell.get("cached_value") if cell.get("formula") else cell.get("value")


def table_sources(sheet):
    regions = [(region, source_cells(region)) for region in sheet.regions]
    all_cells = {address: cell for _, cells in regions for address, cell in cells.items()}
    emitted = []
    for native in sheet.tables:
        left, top, right, bottom = range_boundaries(native.reference)
        bottom -= max(0, native.totals_row_count)
        reference = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
        cells, complete, merged, hidden = _source_range(reference, regions)
        if cells:
            emitted.append(native.reference)
            yield reference, cells, complete, merged, hidden, "native"
    for reference in series_block_ranges(regions, all_cells):
        if any(_overlaps(reference, known) for known in emitted):
            continue
        cells, complete, merged, hidden = _source_range(reference, regions)
        emitted.append(reference)
        yield reference, cells, complete, merged, hidden, "schema"
    for region, _ in regions:
        original = f"{region.start_cell}:{region.end_cell}"
        reference = numeric_body_header(original, all_cells)
        if reference == original or any(_overlaps(reference, known) for known in emitted):
            continue
        cells, complete, merged, hidden = _source_range(reference, regions)
        emitted.append(reference)
        yield reference, cells, complete, merged, hidden, "region"
    grouped = defaultdict(list)
    for schema in sheet.column_schemas:
        grouped[schema.source_range].append(schema)
    for reference in grouped:
        if any(_overlaps(reference, known) for known in emitted):
            continue
        reference = recover_series_header_range(reference, all_cells)
        cells, complete, merged, hidden = _source_range(reference, regions)
        if len(cells) >= 4:
            reference = _trimmed_range(reference, cells, complete)
            emitted.append(reference)
            yield reference, cells, complete, merged, hidden, "schema"
    hidden_rows = {number for region, _ in regions for number in region.hidden_rows}
    merged = {span for region, _ in regions for span in region.merged_ranges}
    for reference in report_ranges(all_cells, hidden_rows, merged):
        reference = recover_series_header_range(reference, all_cells)
        if any(_overlaps(reference, known) for known in emitted):
            continue
        cells, complete, merged, hidden = _source_range(reference, regions)
        if not complete:
            continue  # A sample-inferred bottom is not an original population boundary.
        emitted.append(reference)
        yield reference, cells, complete, merged, hidden, "report"
    for region, cells in regions:
        reference = f"{region.start_cell}:{region.end_cell}"
        if any(_overlaps(reference, known) for known in emitted):
            continue
        segments = _row_segments(reference, cells, region.hidden_rows) if _complete(region) else [(reference, cells)]
        for segment, segment_cells in segments:
            recovered = numeric_body_header(segment, all_cells)
            if recovered != segment:
                segment = recovered
                segment_cells, complete, merged, hidden = _source_range(segment, regions)
                yield segment, segment_cells, complete, merged, hidden, "region"
                continue
            _, top, _, bottom = range_boundaries(segment)
            yield (segment, segment_cells, _complete(region),
                   region.merged_ranges,
                   frozenset(row for row in region.hidden_rows if top <= row <= bottom), "region")


def _overlaps(first, second):
    a, b, c, d = range_boundaries(first)
    e, f, g, h = range_boundaries(second)
    return a <= g and e <= c and b <= h and f <= d


def _trimmed_range(reference, cells, complete=True):
    left, top, right, bottom = range_boundaries(reference)
    top = leaf_header_row(cells, left, top, right, bottom)
    bottom = min(bottom, max(coordinate_to_tuple(key)[0] for key in cells)) if complete else bottom
    return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"


def table_context(sheet, reference, cells):
    left, top, right, _ = range_boundaries(reference)
    result = []
    for number in range(max(1, top - 3), top):
        for column in range(left, right + 1):
            address = f"{get_column_letter(column)}{number}"
            cell = cells.get(address, {})
            value = cell_value(cell)
            if cell.get("bold") and isinstance(value, str) and value.strip():
                result.append((f"{sheet}!{address}", value.strip()[:250]))
    return tuple(result[:12])
