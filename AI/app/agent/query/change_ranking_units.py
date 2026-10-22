"""Inherit a grouped count unit only after complete parent-total reconciliation."""
from dataclasses import replace
from decimal import Decimal

from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries
from app.agent.query.comparison_value_format import people_unit
from app.agent.query.excel_number_formats import parse_number_format
from app.agent.query.metric_observation_cells import _numeric
from app.agent.query.metric_observation_units import observation_unit


def reconciled_group_units(series, index):
    if not series or len({candidate.table_id for candidate, _ in series}) != 1:
        return series
    points = [candidate.observations for candidate, _ in series]
    if len({tuple(point.period for point in items) for items in points}) != 1:
        return series
    sheet, _, reference = series[0][0].table_id.rpartition("!")
    left, top, right, _ = range_boundaries(reference)
    cells = {cell.address: cell for row in index.rows if row.sheet_name == sheet for cell in row.cells}
    selected = {get_column_letter(coordinate_to_tuple(items[0].value.address)[1]) for items in points}
    headers = [cell for cell in cells.values() if coordinate_to_tuple(cell.address)[0] == top
               and left <= coordinate_to_tuple(cell.address)[1] <= right
               and get_column_letter(coordinate_to_tuple(cell.address)[1]) not in selected
               and people_unit("", cell.value)]
    valid = []
    for header in headers:
        parent_column = get_column_letter(coordinate_to_tuple(header.address)[1])
        parents = []
        for offset, point in enumerate(points[0]):
            row = coordinate_to_tuple(point.value.address)[0]
            cell = cells.get(f"{parent_column}{row}")
            if cell is None or not _numeric(cell.value):
                break
            if sum(Decimal(str(items[offset].value.value)) for items in points) != Decimal(str(cell.value)):
                break
            parent = replace(point, label=header, value=cell)
            unit = observation_unit(parent, index.rows)
            if not unit.known or any(_style(items[offset].value) != _style(cell) for items in points):
                break
            parents.append((parent, unit))
        if len(parents) == len(points[0]):
            valid.append(parents)
    if len(valid) != 1:
        return series
    parents = valid[0]
    proof = tuple(dict.fromkeys([parents[0][0].label, *[parent.value for parent, _ in parents],
                                *[point.value for items in points for point in items]]))
    result = []
    for candidate, units in series:
        inferred = tuple(unit if unit.known else replace(parent_unit, references=proof)
                         for unit, (_, parent_unit) in zip(units, parents))
        result.append((candidate, inferred))
    return result


def _style(cell):
    spec = parse_number_format(cell.number_format)
    return spec.prefix, spec.suffix, spec.percent, spec.multiplier, spec.divisor
