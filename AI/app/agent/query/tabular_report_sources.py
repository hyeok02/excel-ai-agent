"""Recover report tables from source geometry, not workbook or business names.

Decorative spacer columns and hidden helper rows often make a single displayed
table appear as dozens of disconnected regions. A strong header row can bridge
those regions, while repeated leading headers keep parallel populations apart.
"""
from collections import Counter, defaultdict

from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries


def report_ranges(cells, hidden_rows, merged=()):
    rows = defaultdict(dict)
    for address, cell in cells.items():
        row, column = coordinate_to_tuple(address)
        value = _value(cell)
        if value is not None and (not isinstance(value, str) or value.strip()):
            rows[row][column] = cell
    occupied = []
    vertical = [range_boundaries(span) for span in merged]
    for number in sorted(rows):
        # Inference requires a source formatting signal. Plain rectangular tables
        # continue through the existing schema/region path instead of guessing that
        # an arbitrary text record is a header.
        header_cells = dict(rows[number])
        for left, top, right, bottom in vertical:
            if left == right and top < number <= bottom and left not in header_cells:
                anchor = rows.get(top, {}).get(left)
                if anchor is not None:
                    header_cells[left] = anchor
        headers = [(column, cell) for column, cell in sorted(header_cells.items())
                   if isinstance(_value(cell), str) and _value(cell).strip()
                   and cell.get("bold")]
        if len(headers) < 2:
            continue
        for group in _header_groups(headers):
            if len(group) < 2:
                continue
            columns = [column for column, _ in group]
            left, right = columns[0], columns[-1]
            if any(a <= left and right <= c and b <= number <= d
                   for a, b, c, d in occupied):
                continue
            first = _first_data_row(number, columns, rows, hidden_rows)
            if first is None:
                continue
            bottom = first
            for following in range(first + 1, max(rows) + 1):
                if following in hidden_rows:
                    bottom = following
                    continue
                selected = [rows.get(following, {}).get(column) for column in columns]
                populated = [cell for cell in selected if cell is not None]
                if not populated or _new_heading(populated, len(columns)):
                    break
                bottom = following
            occupied.append((left, number, right, bottom))
            yield f"{get_column_letter(left)}{number}:{get_column_letter(right)}{bottom}"


def _header_groups(headers):
    gaps = [b[0] - a[0] for a, b in zip(headers, headers[1:])]
    spacing = Counter(gaps).most_common(1)[0][0] if gaps else 1
    groups, current = [], []
    for column, cell in headers:
        value = _value(cell).strip()
        repeat = current and value == _value(current[0][1]).strip()
        distance = column - current[-1][0] if current else 0
        gap = current and (spacing == 1 and distance > 1 or distance > max(2, spacing * 2))
        if current and len(current) >= 2 and (repeat or gap):
            groups.append(current)
            current = []
        current.append((column, cell))
    if current:
        groups.append(current)
    return groups


def _first_data_row(header_row, columns, rows, hidden_rows):
    for number in range(header_row + 1, header_row + 9):
        if number in hidden_rows:
            continue
        selected = [rows.get(number, {}).get(column) for column in columns]
        populated = [cell for cell in selected if cell is not None]
        if not populated:
            continue
        if len(populated) < 2 or _new_heading(populated, len(columns)):
            return None
        return number
    return None


def _new_heading(cells, column_count):
    return (len(cells) == 1 and column_count > 2
            or all(cell.get("bold") and isinstance(_value(cell), str) for cell in cells))


def _value(cell):
    return cell.get("cached_value") if cell.get("formula") else cell.get("value")


def leaf_header_row(cells, left, top, right, bottom):
    """Do not treat a sparse header-group row as observations' leaf labels."""
    def headers(number):
        return [cells.get(f"{get_column_letter(column)}{number}", {})
                for column in range(left, right + 1)
                if _value(cells.get(f"{get_column_letter(column)}{number}", {})) not in (None, "")]
    for number in range(top + 1, min(bottom, top + 3) + 1):
        previous, candidate = headers(top), headers(number)
        if (len(candidate) < len(previous) or len(candidate) < 2
                or not all(cell.get("bold") and isinstance(_value(cell), str) for cell in candidate)):
            break
        top = number
    return top
