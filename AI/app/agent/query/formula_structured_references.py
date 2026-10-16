"""Exact table metadata binding for bounded, rectangular structured references."""
import re

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries

from app.agent.query.formula_resolution_models import FormulaUnavailable

_LEAF = r"\[((?:'[\[\]#@']|[^\[\]])*)\]"
_ITEMS = {"#data", "#all", "#headers", "#totals", "#this row"}


def structured_reference(raw, sheet, resolver):
    marker = raw.find("[")
    name, spec = (raw[:marker], raw[marker:]) if marker >= 0 else (raw, "")
    current = resolver.reference_stack[-1] if resolver.reference_stack else ""
    address = current.rsplit("!", 1)[-1]
    row, column = coordinate_to_tuple(address) if address else (0, 0)
    metadata = resolver.index.reference_metadata
    tables = [table for table in metadata.tables if table.name.casefold() == name.casefold()] if name else [
        table for table in metadata.tables if table.sheet == sheet and _contains(table.reference, row, column)]
    if len(tables) != 1:
        raise FormulaUnavailable("unsupported_reference", (raw,))
    table = tables[0]
    if table.sheet not in metadata.allowed_sheets:
        raise FormulaUnavailable("missing_reference", (raw,))
    left, top, right, bottom = _bounds(table)
    item, first, last = _selectors(spec)
    if first is not None:
        first, last = _column(table, first), _column(table, last)
        if first > last:
            raise FormulaUnavailable("unsupported_reference", (raw,))
        left, right = left + first, left + last
    start, end = top + table.headers, bottom - table.totals
    if item == "#headers":
        if not table.headers:
            raise FormulaUnavailable("unsupported_reference", (raw,))
        end = start = top
    elif item == "#totals":
        if not table.totals:
            raise FormulaUnavailable("unsupported_reference", (raw,))
        end = start = bottom
    elif item == "#all":
        start, end = top, bottom
    elif item == "#this row":
        if table.sheet != sheet or not start <= row <= end:
            raise FormulaUnavailable("unsupported_reference", (raw,))
        start = end = row
    if start > end:
        raise FormulaUnavailable("unsupported_reference", (raw,))
    _formula_inputs(table, left, right, start, end, resolver)
    origin, target = f"{get_column_letter(left)}{start}", f"{get_column_letter(right)}{end}"
    rectangle = origin if origin == target else f"{origin}:{target}"
    return "'" + table.sheet.replace("'", "''") + "'!" + rectangle


def _bounds(table):
    try:
        bounds = range_boundaries(table.reference)
        left, top, right, bottom = bounds
        if (not table.name or not all(table.columns) or None in bounds or left > right or top > bottom or right > 16384 or bottom > 1048576
                or table.headers not in (0, 1) or table.totals not in (0, 1)
                or len(table.columns) != right - left + 1 or len(set(name.casefold() for name in table.columns)) != len(table.columns)):
            raise ValueError()
        return bounds
    except (ValueError, TypeError):
        raise FormulaUnavailable("unsupported_reference", (table.name,)) from None


def _contains(reference, row, column):
    try:
        left, top, right, bottom = range_boundaries(reference)
        return left <= column <= right and top <= row <= bottom
    except (ValueError, TypeError):
        return False


def _selectors(spec):
    if not spec:
        return "#data", None, None
    match = re.fullmatch(_LEAF, spec)
    if match:
        token = match[1]
        if token.casefold() in _ITEMS:
            return token.casefold(), None, None
        return ("#this row", token[1:], token[1:]) if token.startswith("@") else ("#data", token, token)
    match = re.fullmatch(r"\[@" + _LEAF + r"\]", spec)
    if match:
        return "#this row", match[1], match[1]
    match = re.fullmatch(r"\[" + _LEAF + ":" + _LEAF + r"\]", spec)
    if match:
        return "#data", match[1], match[2]
    match = re.fullmatch(r"\[" + _LEAF + "," + _LEAF + "(?:" + ":" + _LEAF + r")?\]", spec)
    if match and match[1].casefold() in _ITEMS:
        return match[1].casefold(), match[2], match[3] or match[2]
    raise FormulaUnavailable("unsupported_reference", (spec,))


def _column(table, raw):
    if re.search(r"[\[\]#@']", re.sub(r"'([\[\]#@'])", "", raw)):
        raise FormulaUnavailable("unsupported_reference", (raw,))
    name = re.sub(r"'([\[\]#@'])", r"\1", raw)
    matches = [i for i, column in enumerate(table.columns) if column.casefold() == name.casefold()]
    if len(matches) != 1:
        raise FormulaUnavailable("unsupported_reference", (raw,))
    return matches[0]


def _formula_inputs(table, left, right, top, bottom, resolver):
    if (right - left + 1) * (bottom - top + 1) > resolver.limits.max_range_cells:
        raise FormulaUnavailable("resource_limit")
    origin, first, _, last = _bounds(table)
    for column in range(left, right + 1):
        position = column - origin
        for row in range(top, bottom + 1):
            calculated = position in table.calculated_columns and first + table.headers <= row <= last - table.totals
            total = position in table.total_columns and table.totals and row == last
            if calculated or total:
                reference = f"{table.sheet}!{get_column_letter(column)}{row}"
                cell = resolver.cells.get(reference)
                if cell is None or cell.value is None and not cell.formula:
                    raise FormulaUnavailable("unsupported_reference", (reference,))
