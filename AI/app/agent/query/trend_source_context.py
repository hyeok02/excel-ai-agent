"""Preserve verified source captions so identical numbers do not merge entities."""
from openpyxl.utils import coordinate_to_tuple, get_column_letter

from app.agent.query.workbook_source import source_for
from app.agent.query.workbook_source_models import SourceReadError


def verified_context(entries, index):
    if not entries:
        return ()
    cells = {cell.reference: cell for row in index.rows for cell in row.cells}
    missing = [(reference, expected) for reference, expected in entries if reference not in cells]
    provider = source_for(index)
    if missing and provider is None:
        return None
    try:
        for sheet in {reference.rpartition("!")[0] for reference, _ in missing}:
            addresses = [reference.rpartition("!")[2] for reference, _ in missing
                         if reference.rpartition("!")[0] == sheet]
            coordinates = [coordinate_to_tuple(address) for address in addresses]
            top, bottom = min(row for row, _ in coordinates), max(row for row, _ in coordinates)
            left, right = min(column for _, column in coordinates), max(column for _, column in coordinates)
            reference = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
            scan = provider.scan_rows(sheet, reference)
            with scan as rows:
                for row in rows:
                    cells.update((cell.reference, cell) for cell in row.cells)
            if not scan.complete:
                return None
    except (SourceReadError, TypeError, ValueError):
        return None
    result = []
    for reference, expected in entries:
        cell = cells.get(reference)
        if cell is None or not isinstance(cell.value, str) or cell.value.strip()[:250] != expected:
            return None
        result.append(cell)
    return tuple(result)
