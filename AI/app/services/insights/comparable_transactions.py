"""Build bounded comparable-transaction facts from a detected table layout."""
import math
from statistics import median

from openpyxl.utils import get_column_letter

from app.services.insights.comparable_layout import comparable_layout, finite, metric_kind


def extract_comparable_transactions(sheet_name, regions):
    layout = comparable_layout(sheet_name, regions)
    if not layout:
        return None
    rows, cells = layout["rows"], layout["cells"]
    subject_cell = rows[layout["subject"]][layout["name_column"]]
    metrics = [
        item for kind in ("ebitda_multiple", "transaction_value")
        if (item := _metric(sheet_name, cells, layout, kind))
    ]
    if not metrics:
        return None
    subject_reference = _reference(sheet_name, subject_cell["cell"])
    for metric in metrics:
        metric["evidence"].insert(1, subject_reference)
    return {
        "subject": str(subject_cell["value"]).strip(),
        "subject_cell": subject_cell["cell"],
        "peer_count": len(layout["peers"]),
        "peer_name_range": _range(sheet_name, layout["name_column"], layout["peers"]),
        "metrics": metrics,
    }


def _metric(sheet, cells, layout, kind):
    column, label_cell = layout["headers"][kind]
    subject = _at(cells, column, layout["subject"])
    peer_cells = [_at(cells, column, row) for row in layout["peers"]]
    valid = [cell for cell in peer_cells if finite(cell)]
    if not finite(subject) or len(valid) < 3:
        return None
    middle = median(float(cell["value"]) for cell in valid)
    if middle == 0:
        return None
    median_cell = _at(cells, column, layout["median_row"])
    evidence = [_reference(sheet, label_cell["cell"]),
                _reference(sheet, subject["cell"]),
                _range(sheet, column, layout["peers"])]
    median_matches = finite(median_cell) and math.isclose(
        float(median_cell["value"]), middle, rel_tol=1e-9, abs_tol=1e-9
    )
    if median_matches:
        evidence.append(_reference(sheet, median_cell["cell"]))
    return {
        "kind": metric_kind(str(label_cell["value"]))[0],
        "label": str(label_cell["value"]).strip(),
        "header_cell": label_cell["cell"],
        "subject_value": float(subject["value"]), "subject_cell": subject["cell"],
        "median": middle, "median_cell": median_cell["cell"] if median_matches else None,
        "difference": float(subject["value"]) - middle,
        "difference_percent": (float(subject["value"]) / middle - 1) * 100,
        "valid_count": len(valid), "peer_count": len(layout["peers"]),
        "peer_range": _range(sheet, column, layout["peers"]), "evidence": evidence,
    }


def _at(cells, column, row):
    return cells.get(f"{get_column_letter(column)}{row}")


def _reference(sheet, cell):
    return f"'{str(sheet).replace(chr(39), chr(39) * 2)}'!{cell}"


def _range(sheet, column, rows):
    letter = get_column_letter(column)
    return _reference(sheet, f"{letter}{min(rows)}:{letter}{max(rows)}")
