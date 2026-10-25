"""Compare a focal deal with its peers: one median per source column or ratio."""
import math
from statistics import median

from openpyxl.utils import get_column_letter

from app.services.insights.facts.comparable_layout import finite

MIN_VALUES = 3


def source_multiples(sheet, layout):
    """One reported multiple per earnings basis: the column with most values."""
    best = {}
    for item in layout["columns"]:
        if item["role"] != "multiple":
            continue
        metric = column_metric(sheet, layout, item, "multiple")
        key = (metric or {}).get("valid_count", -1), item["priority"], -item["column"]
        if metric and (item["detail"] not in best or key > best[item["detail"]][0]):
            best[item["detail"]] = (key, {**metric, "basis": item["detail"]})
    return [metric for _, metric in best.values()]


def derived_multiples(sheet, layout, reported_bases):
    """Deal value ÷ target earnings, only for a basis the source never reports.

    A provider's own NA is not second-guessed: if any column already reports a
    multiple on that basis, its coverage stands as the source states it.
    """
    metrics = []
    for basis, earnings in _target_earnings(layout).items():
        if basis in reported_bases:
            continue
        roles = ("equity",) if basis in {"net_income", "book"} else ("enterprise", "size")
        options = [_ratio(sheet, layout, top, earnings, basis)
                   for top in layout["columns"] if top["role"] in roles]
        options = [item for item in options if item and item["valid_count"] * 2 >= item["peer_count"]]
        if options:
            metrics.append(max(options, key=lambda item: item["valid_count"]))
    return metrics


def column_metric(sheet, layout, item, kind):
    column = item["column"]
    subject = _at(layout, column, layout["subject"])
    peers = [_at(layout, column, row) for row in layout["peers"]]
    values = [float(cell["value"]) for cell in peers if finite(cell)]
    if not finite(subject) or len(values) < MIN_VALUES or (middle := median(values)) == 0:
        return None
    median_cell = _at(layout, column, layout["median_row"])
    matches = finite(median_cell) and math.isclose(float(median_cell["value"]), middle,
                                                   rel_tol=1e-9, abs_tol=1e-9)
    evidence = [_reference(sheet, item["cell"]["cell"]), _reference(sheet, subject["cell"]),
                _range(sheet, column, layout["peers"])]
    if matches:
        evidence.append(_reference(sheet, median_cell["cell"]))
    return {**_comparison(float(subject["value"]), middle, len(values), len(layout["peers"])),
            "kind": kind, "unit": "multiple" if kind == "multiple" else "amount",
            "label": str(item["cell"]["value"]).strip(), "header_cell": item["cell"]["cell"],
            "subject_cell": subject["cell"], "median_cell": median_cell["cell"] if matches else None,
            "peer_range": _range(sheet, column, layout["peers"]), "derived": False,
            "evidence": evidence}


def _ratio(sheet, layout, top, bottom, basis):
    def value(item, row):
        cell = _at(layout, item["column"], row)
        return float(cell["value"]) if finite(cell) else None
    def ratio(row):
        numerator, denominator = value(top, row), value(bottom, row)
        return numerator / denominator if numerator is not None and denominator and denominator > 0 else None
    subject = ratio(layout["subject"])
    values = [item for row in layout["peers"] if (item := ratio(row)) is not None]
    if subject is None or len(values) < MIN_VALUES or (middle := median(values)) == 0:
        return None
    cells = [_at(layout, item["column"], layout["subject"])["cell"] for item in (top, bottom)]
    return {**_comparison(subject, middle, len(values), len(layout["peers"])),
            "kind": "multiple", "unit": "multiple", "basis": basis, "derived": True,
            "numerator_role": top["role"],
            "label": f"{_label(top)} / {_label(bottom)}", "header_cell": top["cell"]["cell"],
            "numerator": {"label": _label(top), "value": value(top, layout["subject"])},
            "denominator": {"label": _label(bottom), "value": value(bottom, layout["subject"])},
            "subject_cell": cells[0], "median_cell": None,
            "evidence": [_reference(sheet, top["cell"]["cell"]), _reference(sheet, bottom["cell"]["cell"]),
                         *(_reference(sheet, cell) for cell in cells),
                         *(_range(sheet, item["column"], layout["peers"]) for item in (top, bottom))]}


def _target_earnings(layout):
    """Target-company earnings; a buyer's or an ambiguous duplicate is not used."""
    found = {}
    for item in layout["columns"]:
        if item["role"] == "earnings" and item["side"] != "buyer":
            found.setdefault(item["detail"], []).append(item)
    targets = {basis: [item for item in items if item["side"] == "target"] or items
               for basis, items in found.items()}
    return {basis: items[0] for basis, items in targets.items() if len(items) == 1}


def _comparison(subject, middle, valid, peers):
    return {"subject_value": subject, "median": middle, "difference": subject - middle,
            "difference_percent": (subject / middle - 1) * 100,
            "valid_count": valid, "peer_count": peers}


def _label(item):
    return " ".join(str(item["cell"]["value"]).split())


def _at(layout, column, row):
    return layout["cells"].get(f"{get_column_letter(column)}{row}")


def _reference(sheet, cell):
    return f"'{str(sheet).replace(chr(39), chr(39) * 2)}'!{cell}"


def _range(sheet, column, rows):
    letter = get_column_letter(column)
    return _reference(sheet, f"{letter}{min(rows)}:{letter}{max(rows)}")
