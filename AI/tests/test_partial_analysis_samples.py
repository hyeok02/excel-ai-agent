"""상한에 걸린 영역이 분석에서 통째로 사라지지 않는지 확인한다."""
from openpyxl import Workbook

from app.services.region_detector import CellRegion
from app.services.workbook_details import analysis_samples
from app.services.workbook_details.analysis_samples import collect_analysis_sample
from app.services.workbook_details.regions import summarize_regions

def test_region_over_the_cap_keeps_the_cells_it_already_read(monkeypatch) -> None:
    monkeypatch.setattr(analysis_samples, "MAX_ANALYSIS_SHEET_CELLS", 100)
    monkeypatch.setattr(analysis_samples, "MAX_ANALYSIS_REGION_CELLS", 100)
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 31):
        for column in range(1, 9):
            sheet.cell(row, column, f"r{row}c{column}")
    try:
        region = summarize_regions(sheet, [CellRegion("A1", "H30", 240)])[0]
        cells = [cell for row in region.analysis_rows for cell in row]

        # 상한을 넘겼다고 읽은 셀까지 버리면 이 영역은 분석에서 통째로 사라진다.
        assert not region.analysis_complete
        assert len(cells) == 100
        assert cells[0]["address"] == "A1"
        assert cells[-1]["address"] == "D13"
    finally:
        workbook.close()


def test_scan_cap_also_keeps_the_cells_it_already_read(monkeypatch) -> None:
    monkeypatch.setattr(analysis_samples, "MAX_ANALYSIS_SCAN_CELLS", 12)
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 11):
        for column in range(1, 9):
            sheet.cell(row, column, f"r{row}c{column}")
    try:
        rows, complete = collect_analysis_sample(sheet, None, (1, 1, 8, 10), 2048)

        assert not complete
        assert [cell["address"] for row in rows for cell in row] == [
            "A1", "B1", "C1", "D1", "E1", "F1", "G1", "H1", "A2", "B2", "C2", "D2",
        ]
    finally:
        workbook.close()
