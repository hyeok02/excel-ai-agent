from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from app.agent.query.index import build_workbook_data_index
from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.tabular_execution import execute_tabular_query
from app.agent.query.tabular_query_models import TabularQueryPlan
from app.services.workbook_parser import parse_workbook


def _parsed(book):
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    return parse_workbook("test.xlsx", content), build_workbook_data_index("test.xlsx", content)


def test_native_table_totals_are_metadata_not_observations():
    book = Workbook()
    sheet = book.active
    for row in [["ID", "Amount"], ["a", 10], ["b", 20], ["TOTAL", 30]]:
        sheet.append(row)
    sheet.add_table(Table(displayName="Measurements", ref="A1:B4", totalsRowCount=1))
    summary, index = _parsed(book)
    assert summary.sheets[0].tables[0].totals_row_count == 1
    table = build_query_tables(summary)[0]
    assert table.reference == "A1:B3" and len(table.rows) == 2
    plan = TabularQueryPlan(operation="numeric_summary", table_id=table.id, column_ids=["B"])
    answer = execute_tabular_query(plan, summary, index, "amounts")
    assert "합계 30" in answer.answer and "평균 15" in answer.answer
    assert "최댓값 20" in answer.answer


def test_hidden_row_does_not_split_the_visible_table_into_certified_fragments():
    book = Workbook()
    sheet = book.active
    for row in [
        ["Address", "Asset", "Phase", "Connection", "Transfer", "Collector", "Notes"],
        ["a", "e1", "ready", "bad", "ready", "bad", "new"],
        ["b", "e2", "ready", "bad", "ready", "bad", "new"],
        ["c", "e3", "ready", "good", "ready", "bad", "new"],
        ["d", "e4", "ready", "good", "ready", "bad", "new"],
    ]:
        sheet.append(row)
    sheet.row_dimensions[3].hidden = True
    summary, index = _parsed(book)
    tables = build_query_tables(summary)
    assert len(tables) == 1
    table = tables[0]
    assert table.reference == "A1:G5"
    assert [row.row_number for row in table.rows] == [2, 4, 5]
    plan = TabularQueryPlan(operation="frequencies", table_id=table.id, column_ids=["G"])
    answer = execute_tabular_query(plan, summary, index, "notes")
    assert "‘new’ 3건" in answer.answer and "숨겨진 1개 행" in answer.answer
    count_plan = plan.model_copy(update={"operation": "row_count", "column_ids": []})
    assert "3건" in execute_tabular_query(count_plan, summary, index, "count").answer


def test_formula_only_row_remains_an_observation_with_unknown_values():
    book = Workbook()
    for row in [["ID", "Amount"], ["a", 10], ["=1", "=2"], ["b", 20]]:
        book.active.append(row)
    summary, index = _parsed(book)
    table = build_query_tables(summary)[0]
    assert [row.row_number for row in table.rows] == [2, 3, 4]
    plan = TabularQueryPlan(operation="frequencies", table_id=table.id, column_ids=["B"])
    assert execute_tabular_query(plan, summary, index, "amount counts") is None


def test_hidden_header_does_not_promote_first_data_record_to_a_header():
    book = Workbook()
    for row in [
        ["Code", "Asset", "Phase", "Connection", "Transfer", "Collector", "Notes"],
        ["a", "e1", "w", "n", "c", "d", "new"],
        ["b", "e2", "w", "n", "c", "d", "new"],
        ["c", "e3", "w", "y", "c", "d", "new"],
        ["d", "e4", "w", "y", "c", "d", "new"],
    ]:
        book.active.append(row)
    book.active.row_dimensions[1].hidden = True
    summary, _ = _parsed(book)
    assert build_query_tables(summary) == []
