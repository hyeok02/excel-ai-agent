"""Advanced computation is bound to uploaded bytes and discovered source IDs."""
import asyncio

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
import pytest

from app.api import workbook_questions as endpoint
from app.agent import create_default_tool_registry
from app.agent.query.conversation_models import ContextResolutionDraft, QuestionConversation
from app.agent.query.service import WorkbookQuestionService
from app.main import app
from tests.support.advanced_qna_fixtures import checked, condition, group_plan, has_group, source
from tests.support.large_workbook_fixtures import (
    FixedClassifier, NoGeneratedAnswer, ask, fresh_index, offline, parsed, plan_for, tiny_caps,
)


@pytest.fixture
def client():
    saved = dict(app.dependency_overrides)
    app.dependency_overrides[endpoint.get_question_answer_generator] = lambda: NoGeneratedAnswer()
    app.dependency_overrides[endpoint.get_question_context_resolver] = lambda: None
    with TestClient(app) as result:
        yield result
    app.dependency_overrides.clear()
    app.dependency_overrides.update(saved)


def test_api_having_rescans_each_uploaded_original_and_never_prior_response(client, tiny_caps, monkeypatch):
    from tests.support.advanced_qna_fixtures import records
    summary, _, original = source()
    changed = records()
    changed[-1][1] = 105
    _, _, other = source(changed)
    plan = group_plan(summary, having=[condition(100)])
    monkeypatch.setattr(endpoint, "_intent_classifier", FixedClassifier(plan))
    for content, has_alpha in [(original, False), (other, True), (original, False)]:
        response = client.post("/api/v1/workbooks/questions",
                               data={"question": "Records 시트 Department별 Amount 합계가 100 초과인 그룹 보여줘"},
                               files={"file": ("same-source-name.xlsx", content)})
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["status"] == "answered", body
        assert body["verification"]["level"] == "request_checked", body
        assert "full_source_scan" in body["verification"]["checks"], body
        assert has_group(body["answer"], "Alpha") == has_alpha, body
        if has_alpha:
            assert "합계 175" in body["answer"]


@pytest.mark.parametrize("reordered", [False, True])
@pytest.mark.parametrize("feature", ["having", "distinct", "partition"])
def test_renamed_reordered_headers_and_moved_native_table_work(tiny_caps, reordered, feature):
    headers = ["Batch", "Volume", "Zone", "Party"]
    data = [["First", 40, "West", "Pone"], ["Second", 20, "West", "Ptwo"],
            ["Third", 30, "East", "Pone"], ["Fourth", 10, "East", "Pthree"]]
    positions = [2, 0, 3, 1] if reordered else list(range(4))
    matrix = [[row[position] for position in positions] for row in [headers, *data]]
    book = Workbook()
    sheet = book.active
    sheet.title = "Records"
    for row_number, row in enumerate(matrix, 4):
        for column_number, value in enumerate(row, 4):
            sheet.cell(row_number, column_number, value)
    sheet.add_table(Table(displayName="ArbitrarySource", ref="D4:G8"))
    summary, index, _ = parsed(book, filename="unrelated-workbook.xlsx")
    ids = {header: get_column_letter(4 + matrix[0].index(header)) for header in headers}
    if feature == "having":
        plan = plan_for(summary, columns=(ids["Volume"],), group_by=[ids["Zone"]],
                        having=[condition(50, column=ids["Volume"])])
        question = "Records 시트 Zone별 Volume 합계가 50 초과인 그룹 보여줘"
    elif feature == "distinct":
        plan = plan_for(summary, operation="distinct_count", columns=(ids["Party"],))
        question = "Records 시트 Party 중복 제외 개수 알려줘"
    else:
        plan = plan_for(summary, operation="row_list", columns=(ids["Batch"], ids["Volume"]),
                        group_by=[ids["Zone"]], order_by=ids["Volume"], per_group_limit=1)
        question = "Records 시트 각 Zone별 Volume 상위 1개 Batch와 Volume 보여줘"
    answer = ask(question, summary, index, plan)
    checked(answer, scan=True)
    if feature == "having":
        assert "합계 60" in answer.answer and not has_group(answer.answer, "East", "Zone")
    elif feature == "distinct":
        assert "3건" in answer.answer or "3개" in answer.answer
    else:
        assert "First" in answer.answer and "Third" in answer.answer
        assert "Second" not in answer.answer and "Fourth" not in answer.answer


def test_followup_share_rebuilds_original_groups_instead_of_prior_having_result(tiny_caps):
    summary, index, content = source()
    first_question = "Records 시트 Department별 Amount 합계가 70 초과인 그룹 보여줘"
    first = ask(first_question, summary, index, group_plan(summary, having=[condition(70)]))
    checked(first, scan=True)
    assert not has_group(first.answer, "Gamma")
    resolved = "Records 시트 Department별 Amount 합계의 전체 비중 알려줘"

    class Resolver:
        async def resolve(self, question, context):
            return ContextResolutionDraft(status="resolved", question=resolved, context_ids=["turn:0"])

    service = WorkbookQuestionService(NoGeneratedAnswer(), create_default_tool_registry(),
                                     classifier=FixedClassifier(group_plan(summary, share_of_total=True)),
                                     context_resolver=Resolver())
    context = QuestionConversation(turns=[{
        "question": first_question,
        "evidence": [f"{item.sheet_name}!{item.reference}" for item in first.evidence],
    }])
    answer = asyncio.run(service.answer("그럼 전체 Department의 전체 비중은?", summary,
                                       fresh_index(summary, content), context))
    assert answer.resolved_question == resolved, answer.model_dump()
    assert has_group(answer.answer, "Gamma"), answer.model_dump()
    assert "분자 25 / 분모 190" in answer.answer
    assert {"context_resolved", "full_source_scan"} <= set(answer.verification.checks)
    assert answer.verification.level == "source_checked"
