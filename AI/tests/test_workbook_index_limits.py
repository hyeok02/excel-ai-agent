from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from app.agent.query import index


def test_formatting_beyond_row_limit_does_not_mark_index_truncated(monkeypatch) -> None:
    monkeypatch.setattr(index, "MAX_ROWS_PER_SHEET", 1)
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "실제 데이터"
    sheet["A3"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")

    result = index.build_workbook_data_index("formatted.xlsx", _content(workbook))

    assert result.truncated is False
    assert [row.row_number for row in result.rows] == [1]


def test_actual_data_beyond_row_limit_remains_truncated(monkeypatch) -> None:
    monkeypatch.setattr(index, "MAX_ROWS_PER_SHEET", 1)
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "첫 행"
    sheet["A3"] = "제한 이후 실제 데이터"

    result = index.build_workbook_data_index("oversized.xlsx", _content(workbook))

    assert result.truncated is True
    assert [row.row_number for row in result.rows] == [1]


def _content(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
