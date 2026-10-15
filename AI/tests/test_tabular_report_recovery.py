from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.formula import ArrayFormula

from app.agent.query.cell_values import safe_indexed_value
from app.agent.query.index import build_workbook_data_index
from app.agent.query.tabular_catalog import build_query_tables, query_catalog
from app.agent.query.tabular_validation import validated_table
from app.services.workbook_details.analysis_samples import collect_analysis_sample
from app.services.workbook_details.snapshots import snapshot_range
from app.services.workbook_parser import parse_workbook


def _parsed(book):
    stream = BytesIO()
    book.save(stream)
    content = stream.getvalue()
    return parse_workbook("renamed.xlsx", content), build_workbook_data_index("renamed.xlsx", content)


def _headers(sheet, row, start, labels, spacing=1):
    for offset, label in enumerate(labels):
        cell = sheet.cell(row, start + spacing * offset, label)
        cell.font = Font(bold=True)


@pytest.mark.parametrize("start_row,start_column", [(4, 2), (12, 7)])
def test_parallel_populations_under_titles_remain_separate_after_relocation(start_row, start_column):
    book = Workbook()
    sheet = book.active
    sheet.title = "Different report"
    for shift, title, values in [(0, "Primary population", [4, 5, 6]),
                                  (5, "Secondary population", [90, 80, 70])]:
        start = start_column + shift
        _headers(sheet, start_row, start, [title])
        _headers(sheet, start_row + 1, start, ["Position", "Entity", "Percent"])
        for offset, value in enumerate(values, 2):
            for column, item in enumerate([offset - 1, f"Unit {offset}", value], start):
                sheet.cell(start_row + offset, column, item)
    summary, index = _parsed(book)
    tables = build_query_tables(summary)
    assert len(tables) == 2
    assert [len(table.rows) for table in tables] == [3, 3]
    assert [table.rows[0].values[table.columns[-1].id] for table in tables] == [4, 90]
    assert all(validated_table(table, index, {column.id for column in table.columns}) for table in tables)
    assert any(text == "Primary population" for _, text in tables[0].context)
    assert any(text == "Secondary population" for _, text in tables[1].context)


def test_spacer_columns_and_hidden_helper_rows_do_not_hide_the_real_report_table():
    book = Workbook()
    sheet = book.active
    sheet.title = "Offering ledger"
    _headers(sheet, 8, 3, ["Entity", "Amount ($000)", "Eligible", "Source ID"], spacing=2)
    sheet.row_dimensions[9].hidden = True
    for column, value in [(3, "REMOTE_ENTITY"), (5, "REMOTE_AMOUNT"), (7, "REMOTE_FLAG")]:
        sheet.cell(9, column, value)
    sheet.column_dimensions["I"].hidden = True
    for row, values in enumerate([["A", 100, "Yes", 17], ["B", 400, "No", 23]], 12):
        for offset, value in enumerate(values):
            sheet.cell(row, 3 + offset * 2, value)
    summary, index = _parsed(book)
    table = next(table for table in build_query_tables(summary) if table.reference == "C8:I13")
    assert [row.row_number for row in table.rows] == [12, 13]
    assert table.hidden_columns == frozenset({"I"})
    assert table.rows[0].values["I"] == 17
    assert validated_table(table, index, {"E", "I"})
    catalog = next(entry for entry in query_catalog(summary, "Source ID") if entry["table_id"] == table.id)
    assert catalog["columns"][-1]["hidden"] is True
    assert all(cell["address"] != "I12" for region in summary.sheets[0].regions
               for row in region.analysis_rows for cell in row)


def test_duplicate_metric_labels_keep_distinct_source_column_ids():
    book = Workbook()
    sheet = book.active
    _headers(sheet, 2, 1, ["Entity", "Amount", "Amount"])
    sheet.append(["A", 12, 120])
    sheet.append(["B", 24, 240])
    summary, index = _parsed(book)
    table = build_query_tables(summary)[0]
    assert [column.id for column in table.columns] == ["A", "B", "C"]
    assert validated_table(table, index, {"B", "C"})


def test_horizontal_presentation_merges_do_not_join_real_data_columns():
    book = Workbook()
    sheet = book.active
    _headers(sheet, 3, 2, ["Entity", "Amount"], spacing=2)
    sheet.merge_cells("B3:C3")
    sheet.merge_cells("D3:E3")
    for row, values in [(4, ["A", 10]), (5, ["B", 20])]:
        sheet.cell(row, 2, values[0])
        sheet.cell(row, 4, values[1])
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
    summary, index = _parsed(book)
    table = build_query_tables(summary)[0]
    assert [column.id for column in table.columns] == ["B", "D"]
    assert validated_table(table, index, {"D"})


def test_empty_hyperlink_metadata_is_not_materialized_as_a_table_attribute():
    book = Workbook()
    sheet = book.active
    _headers(sheet, 3, 2, ["Entity", "Amount"])
    sheet["D3"].hyperlink = "mailto:example@example.test"
    sheet["D3"] = None
    sheet["D3"].font = Font(bold=True)
    for row, values in [(4, ["A", 10]), (5, ["B", 20])]:
        sheet.cell(row, 2, values[0])
        sheet.cell(row, 3, values[1])
    summary, index = _parsed(book)
    table = build_query_tables(summary)[0]
    assert [column.id for column in table.columns] == ["B", "C"]
    assert validated_table(table, index, {"C"})


@pytest.mark.parametrize("cached", ["Column name", 713.25, datetime(2025, 7, 1)])
def test_array_formula_samples_and_snapshots_use_cached_values_not_object_repr(cached):
    formulas, values = Workbook(), Workbook()
    formulas.active["C4"] = ArrayFormula("C4:C5", "=REMOTE_VALUES()")
    values.active["C4"] = cached
    rows, complete = collect_analysis_sample(formulas.active, values.active, (3, 4, 3, 4), 20)
    snapshot = snapshot_range(formulas.active, values.active, 4, 4, 3, 3)[0][0]
    expected = cached.isoformat() if isinstance(cached, datetime) else cached
    assert complete
    assert rows[0][0]["formula"] == snapshot.formula == "=REMOTE_VALUES()"
    assert rows[0][0]["cached_value"] == snapshot.cached_value == expected
    assert rows[0][0]["value"] is None
    assert safe_indexed_value(cached, "=REMOTE_VALUES()") == (expected, "formula_datetime" if isinstance(cached, datetime) else "formula")
