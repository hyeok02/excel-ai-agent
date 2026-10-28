"""Upload-path acceptance for missing-plan recovery, never live-model accuracy.

Only planning is stubbed. Workbook parsing, catalog construction, request scope,
calculation, and evidence verification run through the production HTTP API.
"""
import asyncio
from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table
import pytest

from app.agent.query.intent_models import QuestionIntent
from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.tabular_query_models import TabularFilter, TabularQueryPlan
from evals.capability_cases import workbook
from evals.offline_api import offline_client


@pytest.fixture(params=[
    (4, 2, "Ledger entries", "Account", "Territory", "Net Amount", "Orion", "North", 17, 33),
    (9, 4, "Renamed worksheet", "Project", "Branch", "Quantity", "Maple", "West", 31, 41),
])
def source(request):
    row, column, sheet, entity, category, measure, name, region, first, last = request.param
    matrix = [
        [entity, category, measure], [name, region, first],
        ["Other", "Elsewhere", 90], ["Third", region, last],
    ]
    book = Workbook()
    page = book.active
    page.title = sheet
    values = {}
    for row_offset, cells in enumerate(matrix):
        for col_offset, value in enumerate(cells):
            cell = page.cell(row + row_offset, column + col_offset, value)
            values[cell.coordinate] = value
    end = f"{get_column_letter(column + 2)}{row + 3}"
    page.add_table(Table(displayName="Entries", ref=f"{get_column_letter(column)}{row}:{end}"))
    output = BytesIO()
    book.save(output)
    book.close()
    return dict(content=output.getvalue(), sheet=sheet, entity=entity, category=category,
                measure=measure, name=name, region=region, first=first, total=first + last,
                values=values, row=row, column=column)


class MissingPlan:
    """Fixed proposal from actual source catalog IDs, with no supplied answer."""

    def __init__(self, source, operation):
        self.source, self.operation = source, operation
        self.questions, self.repairs, self.initial_ids = [], [], None

    async def classify(self, question, summary):
        self.questions.append(question)
        self.initial_ids = {table.id for table in build_query_tables(summary)}
        return QuestionIntent(False, "other", from_model=True)

    async def replan(self, question, summary, feedback):
        self.repairs.append((question, feedback))
        tables = build_query_tables(summary)
        assert {table.id for table in tables} == self.initial_ids
        [table] = [table for table in tables if table.sheet_name == self.source["sheet"]]
        columns = {column.header: column.id for column in table.columns}
        aggregate = self.operation == "aggregate"
        filter_header = self.source["category" if aggregate else "entity"]
        filter_value = self.source["region" if aggregate else "name"]
        plan = TabularQueryPlan(
            operation="numeric_summary" if aggregate else "lookup", table_id=table.id,
            column_ids=[columns[self.source["measure"]]],
            statistic="sum" if aggregate else "all",
            filters=[TabularFilter(column_id=columns[filter_header], operator="equals", value=filter_value)],
        )
        return QuestionIntent(False, "aggregation" if aggregate else "cell_lookup",
                              from_model=True, tabular_query=plan)


@pytest.mark.parametrize("operation", ["aggregate", "lookup"])
def test_uploaded_workbook_recovers_missing_plan_using_parsed_source_ids(source, operation):
    template = ("{category}가 {region}인 {measure}의 합계를 부탁해" if operation == "aggregate"
                else "{entity}가 {name}인 {measure}의 값을 부탁해")
    question = template.format(**source)
    planner = MissingPlan(source, operation)

    async def request():
        async with offline_client(planner) as client:
            return await client.post(
                "/api/v1/workbooks/questions", data={"question": question},
                files={"file": ("unrelated-upload-name.xlsx", source["content"])},
            )

    response = asyncio.run(request())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "answered", body
    assert body["verification"]["level"] == "request_checked", body
    assert {"request_alignment", "plan_replanned"} <= set(body["verification"]["checks"])
    expected = source["total"] if operation == "aggregate" else source["first"]
    assert str(expected) in body["answer"], body
    assert planner.questions == [question] and len(planner.repairs) == 1
    assert planner.repairs[0] == (question, {
        "code": "missing_query_plan", "kind": "other", "previous_plan": None,
    })
    evidence = body["evidence"]
    assert evidence and {item["sheet_name"] for item in evidence} == {source["sheet"]}
    values = source["values"]
    covered = set()
    for item in evidence:
        if item["kind"] == "range":
            left, top, right, bottom = range_boundaries(item["reference"])
            addresses = {f"{get_column_letter(column)}{row}"
                         for column in range(left, right + 1) for row in range(top, bottom + 1)}
            assert addresses <= values.keys(), item
            covered.update(addresses)
        else:
            assert item["reference"] in values, item
            assert item["value"] == values[item["reference"]], item
            covered.add(item["reference"])
    measure_column = get_column_letter(source["column"] + 2)
    required = {f"{measure_column}{source['row'] + 1}"}
    if operation == "aggregate":
        required.add(f"{measure_column}{source['row'] + 3}")
    assert required <= covered


def test_uploaded_ambiguous_source_does_not_retry_to_arbitrarily_verified_scope():
    content = workbook({
        "North": [["Record", "Amount"], ["Entry Alpha", 10], ["Entry Beta", 20]],
        "South": [["Record", "Amount"], ["Entry Gamma", 70], ["Entry Delta", 80]],
    })
    planner = MissingPlan({}, "aggregate")

    async def request():
        async with offline_client(planner) as client:
            return await client.post(
                "/api/v1/workbooks/questions", data={"question": "Amount 합계는 얼마야?"},
                files={"file": ("same-upload-name.xlsx", content)},
            )

    response = asyncio.run(request())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "insufficient_evidence", body
    assert body["verification"]["level"] == "unverified", body
    assert body["clarification_options"], body
    assert "plan_replanned" not in body["verification"]["checks"]
    assert planner.repairs == []
