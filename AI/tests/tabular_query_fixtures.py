from io import BytesIO

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.agent.query.index import build_workbook_data_index
from app.services.workbook_details.models import RegionSummary
from app.services.workbook_parsing.models import SheetSummary, WorkbookSummary


def table_fixture(blocks, sheet_name="Arbitrary", complete=True):
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    regions = []
    for start_row, start_column, matrix in blocks:
        source = []
        for row_number, values in enumerate(matrix, start_row):
            row = []
            for number, value in enumerate(values, start_column):
                address = f"{get_column_letter(number)}{row_number}"
                sheet.cell(row_number, number, value)
                if value is not None:
                    row.append({"address": address, "value": value})
            if row:
                source.append(row)
        start = f"{get_column_letter(start_column)}{start_row}"
        end = f"{get_column_letter(start_column + len(matrix[0]) - 1)}{start_row + len(matrix) - 1}"
        regions.append(RegionSummary(
            start_cell=start, end_cell=end, cell_count=sum(map(len, source)),
            title=None, row_count=len(matrix), column_count=len(matrix[0]),
            merged_ranges=[], header_paths=[], preview_rows=[], is_truncated=True,
            analysis_rows=source, analysis_complete=complete,
        ))
    stream = BytesIO()
    book.save(stream)
    summary = WorkbookSummary("arbitrary.xlsx", 1, [SheetSummary(
        name=sheet_name, rows=sheet.max_row, columns=sheet.max_column,
        formula_count=0, table_count=0, chart_count=0, formulas=[],
        region_count=len(regions), regions=regions,
    )])
    return summary, build_workbook_data_index("arbitrary.xlsx", stream.getvalue())
