from io import BytesIO

from openpyxl import Workbook
import pytest

from app.agent.query.index import build_workbook_data_index
from app.agent.writeback.candidates import select_writeback_candidates
from tests.support.writeback_target_fixtures import proposal


def _workbook_with_a_wide_leading_sheet() -> bytes:
    """
    앞에 넓은 시트, 뒤에 정답 셀이 있는 워크북.

    실제 사고가 난 워크북의 구조를 줄여서 재현한다. 앞 시트가 후보 수 상한을
    다 써버리면 뒤 시트의 정답 셀이 모델에게 전달되지 않는다.
    """
    workbook = Workbook()
    wide = workbook.active
    wide.title = "Overview"
    wide.append([f"Metric {column}" for column in range(40)])
    for row in range(60):
        # 요청에 적을 값(1081)과 겹치지 않는 범위로 채운다.
        wide.append([50_000 + row * 40 + column for column in range(40)])

    detail = workbook.create_sheet("Headcount")
    detail.append(["Department", "Services", "Support"])
    detail.append(["2025-06-01", 1081, 7])
    detail.append(["2023-09-01", 1277, 11])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _references(instruction: str) -> set[str]:
    content = _workbook_with_a_wide_leading_sheet()
    index = build_workbook_data_index("headcount.xlsx", content)
    return {cell.reference for cell in select_writeback_candidates(instruction, index)}


def test_cell_holding_the_stated_value_reaches_the_model() -> None:
    references = _references("Services 인원을 1081에서 981로 줄여줘")

    assert "Headcount!B2" in references


def test_wide_leading_sheet_does_not_crowd_out_the_stated_value() -> None:
    references = _references("Services 인원을 1081에서 981로 줄여줘")
    overview_cells = {item for item in references if item.startswith("Overview!")}

    assert "Headcount!B2" in references
    assert len(overview_cells) < len(references)


def test_selection_still_works_without_a_stated_value() -> None:
    references = _references("Headcount 시트 Services 값을 981로 바꿔줘")

    assert any(item.startswith("Headcount!") for item in references)


def test_explicit_anchored_address_with_korean_particle_stays_a_priority():
    references = _references("Headcount!$b$3을 999로 바꿔줘")
    assert "Headcount!B3" in references


def test_explicit_range_prioritizes_all_interior_cells():
    references = _references("Headcount!$b$2:$c$3을 999로 바꿔줘")
    assert {"Headcount!B2", "Headcount!C2", "Headcount!B3", "Headcount!C3"} <= references


def _replacement_collision_workbook(sheet_name, period, metric, replacement):
    book = Workbook()
    wide = book.active
    wide.title = "Unrelated Overview"
    wide.append([f"Measure {column}" for column in range(40)])
    for _ in range(20):
        wide.append([replacement] * 40)
    source = book.create_sheet(sheet_name)
    source.append(["Period", metric, "Other measure"])
    source.append([period, 37, 18])
    source.append(["Other period", 17, replacement])
    buffer = BytesIO()
    book.save(buffer)
    book.close()
    return buffer.getvalue()


@pytest.mark.parametrize("sheet_name,period,metric,replacement", [
    ("Measured Volumes", "2008 Jan", "Predominantly food stores", 38),
    ("실험_집계", "2024 하반기", "납품 건수", 912),
])
def test_replacement_values_cannot_crowd_source_labels_out_of_candidate_budget(sheet_name, period, metric, replacement):
    content = _replacement_collision_workbook(sheet_name, period, metric, replacement)
    instruction = f"'{sheet_name}' {period} {metric} 값을 {replacement}로 수정해줘"
    index = build_workbook_data_index("unrelated-file-name.xlsx", content)
    selected = select_writeback_candidates(instruction, index)
    assert f"{sheet_name}!B2" in {cell.reference for cell in selected}
    body = proposal(instruction, [(sheet_name, "B2", replacement)], content)
    assert body["status"] == "ready", body
    for reference, value in [("B3", replacement), ("C2", replacement), ("B2", replacement + 1)]:
        wrong = proposal(instruction, [(sheet_name, reference, value)], content)
        assert wrong["status"] == "blocked", wrong


def test_new_text_labels_are_not_source_retrieval_anchors():
    content = _replacement_collision_workbook("Catalog", "Unique period", "Caption", "Replacement")
    index = build_workbook_data_index("text-labels.xlsx", content)
    instruction = "Catalog Unique period Caption 값을 'Replacement'로 수정해줘"
    selected = select_writeback_candidates(instruction, index)
    assert "Catalog!B2" in {cell.reference for cell in selected}
    assert all(cell.sheet_name == "Catalog" for cell in selected)
