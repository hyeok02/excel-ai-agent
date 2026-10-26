from io import BytesIO
from itertools import permutations

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from app.agent.writeback.editor import apply_writeback
from app.agent.writeback.models import WritebackChange
from tests.support.writeback_target_fixtures import proposal


@pytest.mark.parametrize("columns", list(permutations(["Region", "Product", "Amount"])))
@pytest.mark.parametrize("reverse", [False, True])
def test_target_and_preservation_survive_column_row_and_sheet_reordering(columns, reverse):
    book = Workbook()
    book.active.title = "O'Brien Data"
    rows = [{"Region": "Seoul", "Product": "Notebook", "Amount": 10},
            {"Region": "Seoul", "Product": "Monitor", "Amount": 20},
            {"Region": "Busan", "Product": "Notebook", "Amount": 30}]
    if reverse:
        rows.reverse()
    sheet = book.active
    sheet.append(columns)
    for row in rows:
        sheet.append([row[column] for column in columns])
    extra = book.create_sheet("Unrelated", 0)
    extra.append(["Region", "Product", "Amount"])
    extra.append(["Jeju", "Keyboard", 100])
    sheet.merge_cells("A7:C7")
    sheet["A7"] = "Do not change"
    sheet["A9"] = "=SUM(1,2)"
    output = BytesIO()
    book.save(output)
    book.close()
    reference = f"{get_column_letter(columns.index('Amount') + 1)}{4 if reverse else 2}"
    source = output.getvalue()
    body = proposal("Seoul Notebook Amount 값을 99로 수정", [("O'Brien Data", reference, 99)], source)
    assert body["status"] == "ready", body
    assert len(body["changes"]) == 1
    modified, manifest = apply_writeback("generic.xlsx", source,
                                         [WritebackChange.model_validate(body["changes"][0])])
    assert manifest.verified
    original = load_workbook(BytesIO(source))
    result = load_workbook(BytesIO(modified))
    try:
        for before_sheet in original:
            for row in before_sheet:
                for before_cell in row:
                    actual = result[before_sheet.title][before_cell.coordinate]
                    expected = 99 if (before_sheet.title, before_cell.coordinate) == ("O'Brien Data", reference) else before_cell.value
                    assert actual.value == expected
                    assert actual.style_id == before_cell.style_id
    finally:
        original.close()
        result.close()
