from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.agent.writeback.models import WritebackChangeDraft, WritebackProposalDraft
from app.api.workbook_writebacks import get_writeback_generator
from app.main import app
from tests.support.workbook_api_fixtures import upload


def target_workbook():
    book = Workbook()
    book.active.title = "Sales One"
    for name in ("Sales One", "Other", "O'Brien"):
        sheet = book[name] if name in book.sheetnames else book.create_sheet(name)
        sheet.append(["Product", "January", "February"])
        sheet.append(["Notebook", 10, 10])
        sheet.append(["Monitor", 10, 10])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def proposal(instruction, drafts, content=None):
    class Stub:
        async def generate(self, instruction, filename, context):
            return WritebackProposalDraft(summary="대상 검증 대역", changes=[
                WritebackChangeDraft(sheet_name=sheet, reference=ref, new_value=value, reason="검증")
                for sheet, ref, value in drafts
            ])
    app.dependency_overrides[get_writeback_generator] = Stub
    try:
        response = TestClient(app).post(
            "/api/v1/workbooks/writeback-proposals", data={"instruction": instruction},
            files=upload("generic.xlsx", content or target_workbook()),
        )
        assert response.status_code == 200
        return response.json()
    finally:
        app.dependency_overrides.clear()
