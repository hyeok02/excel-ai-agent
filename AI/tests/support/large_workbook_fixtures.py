"""Synthetic workbooks and offline-only service doubles for source rescans."""
import asyncio
from io import BytesIO
import socket

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table
import pytest

from app.agent import create_default_tool_registry
from app.agent.query import index as index_module
from app.agent.query.intent_models import QuestionIntent
from app.agent.query.question_index import build_question_data_index
from app.agent.query.service import WorkbookQuestionService
from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.tabular_query_models import TabularFilter, TabularQueryPlan
from app.services.workbook_details import analysis_samples
from app.services.workbook_parser import parse_workbook


class ForbiddenExternalCall(BaseException):
    pass


@pytest.fixture(autouse=True)
def offline(monkeypatch):
    def blocked(*args, **kwargs):
        raise ForbiddenExternalCall("No network or paid model calls in large-workbook tests")
    monkeypatch.setattr(socket.socket, "connect", blocked)
    monkeypatch.setattr(socket.socket, "connect_ex", blocked)


@pytest.fixture
def tiny_caps(monkeypatch):
    monkeypatch.setattr(index_module, "MAX_INDEXED_CELLS", 24)
    monkeypatch.setattr(index_module, "MAX_ROWS_PER_SHEET", 3)
    monkeypatch.setattr(analysis_samples, "MAX_ANALYSIS_REGION_CELLS", 12)
    monkeypatch.setattr(analysis_samples, "MAX_ANALYSIS_SHEET_CELLS", 12)


def workbook(sheets):
    book = Workbook()
    book.remove(book.active)
    for number, (name, matrix) in enumerate(sheets.items(), 1):
        sheet = book.create_sheet(name)
        for row in matrix:
            sheet.append(row)
        last = f"{get_column_letter(len(matrix[0]))}{len(matrix)}"
        sheet.add_table(Table(displayName=f"SourceTable{number}", ref=f"A1:{last}"))
    return book


def parsed(book, filename="generated-source.xlsx", limits=None):
    stream = BytesIO()
    book.save(stream)
    content = stream.getvalue()
    summary = parse_workbook(filename, content)
    return summary, fresh_index(summary, content, limits), content


def fresh_index(summary, content, limits=None):
    return build_question_data_index(summary.filename, content, summary, limits=limits)


def plan_for(summary, *, sheet="Records", operation="numeric_summary", columns=("B",),
             statistic="sum", filters=(), **kwargs):
    tables = [table for table in build_query_tables(summary) if table.sheet_name == sheet]
    assert len(tables) == 1
    return TabularQueryPlan(operation=operation, table_id=tables[0].id,
                           column_ids=list(columns), statistic=statistic,
                           filters=list(filters), **kwargs)


def equal(column, value):
    return TabularFilter(column_id=column, operator="equals", value=value)


class NoGeneratedAnswer:
    async def generate(self, *args):
        raise ForbiddenExternalCall("Source-rescan tests may not fall back to an answer model")


class FixedClassifier:
    def __init__(self, plan):
        self.plan = plan

    async def classify(self, *args):
        kind = {"lookup": "cell_lookup", "row_list": "row_list", "row_count": "row_count"}.get(
            self.plan.operation, "aggregation")
        return QuestionIntent(False, kind, (), True, self.plan)


def ask(question, summary, index, plan):
    service = WorkbookQuestionService(NoGeneratedAnswer(), create_default_tool_registry(),
                                      classifier=FixedClassifier(plan))
    return asyncio.run(service.answer(question, summary, index))


def assert_stream_checked(answer):
    assert answer.status == "answered", answer.model_dump()
    assert answer.verification.level == "request_checked", answer.model_dump()
    assert "full_source_scan" in answer.verification.checks
    assert "stream_original_table" in answer.selected_tools


def assert_not_certified(answer):
    assert answer.verification.level != "request_checked", answer.model_dump()
    assert answer.status != "answered", answer.model_dump()
