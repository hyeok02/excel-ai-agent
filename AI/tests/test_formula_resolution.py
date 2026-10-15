from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.agent.query.formula_resolution import FormulaResolver
from app.agent.query.index import build_workbook_data_index


def workbook_index(rows, other_rows=()):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    for row in rows:
        sheet.append(row)
    if other_rows:
        other = workbook.create_sheet("Other sheet")
        for row in other_rows:
            other.append(row)
    output = BytesIO()
    workbook.save(output)
    return build_workbook_data_index("arbitrary.xlsx", output.getvalue())


@pytest.mark.parametrize("formula, expected", [
    ("=A1+A2*2", 50), ("=(A1+A2)*2", 60), ("=A1/A2", .5),
    ("=A1^2", 100), ("=-A1^2", 100), ("=20%*A1", 2),
    ("=SUM(A1:A3)", 30), ("=AVERAGE(A1:A3)", 15),
    ("=MIN(A1:A3)", 10), ("=MAX(A1:A3)", 20),
    ("=COUNT(A1:A3)", 2), ("=COUNTA(A1:A3)", 2),
    ("=ROUND(1.255,2)", 1.26), ("=ROUNDUP(-1.21,1)", -1.3),
    ("=ROUNDDOWN(-1.29,1)", -1.2), ("=ABS(-7)", 7),
    ('=IF(A1<A2,"yes","no")', "yes"), ("=IFERROR(1/0,9)", 9),
    ("=AND(A1>0,A2>0)", True), ("=NOT(A1>A2)", True),
    ('="value "&A1', "value 10"), ('=IFERROR("2"*3,0)', 6),
])
def test_deterministic_formulas(formula, expected):
    index = workbook_index([[10, formula], [20], [None]])
    original = next(cell for row in index.rows for cell in row.cells if cell.address == "B1")
    result = FormulaResolver(index).resolve("Data!B1")
    assert result.status == "recovered"
    assert result.value == expected
    assert original.value is None and original.formula == formula


def test_cross_sheet_chain_preserves_original_lineage():
    index = workbook_index([[5, "='Other sheet'!$B$1+A1"]], [[12, "=A1*2"]])
    result = FormulaResolver(index).resolve("Data!B1")
    assert result.value == 29
    assert set(result.inputs) == {"Data!A1", "Other sheet!A1", "Other sheet!B1"}


def test_recovered_formula_dates_and_percent_keep_source_formats():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = date(2026, 5, 3)
    sheet["B1"] = "=A1+2"
    sheet["B1"].number_format = "yyyy-mm-dd"
    sheet["C1"] = "=YEAR(A1)"
    sheet["D1"] = "=DATE(2026,5,4)"
    sheet["D1"].number_format = "yyyy-mm-dd"
    sheet["E1"] = "=25%"
    sheet["E1"].number_format = "0%"
    sheet["F1"] = "=A1=DATE(2026,5,3)"
    sheet["G1"] = "=DATE(2026,5,5)-A1"
    output = BytesIO()
    workbook.save(output)
    index = build_workbook_data_index("dates.xlsx", output.getvalue())
    results = {item.reference: item for item in FormulaResolver(index).resolve_all()}
    assert results["Data!B1"].value.startswith("2026-05-05")
    assert results["Data!C1"].value == 2026
    assert results["Data!D1"].value == "2026-05-04"
    assert results["Data!E1"].value == .25
    assert results["Data!F1"].value is True
    assert results["Data!G1"].value == 2


@pytest.mark.parametrize("formula, expected", [
    ("=SUM(A1:A3)", 10), ("=COUNT(A1:A3)", 1), ("=COUNTA(A1:A3)", 3),
    ('=SUM("2",TRUE,3)', 6), ('=COUNT("2",TRUE,3)', 3),
])
def test_aggregate_reference_text_and_boolean_semantics(formula, expected):
    index = workbook_index([[10, formula], [True], ["20"]])
    result = FormulaResolver(index).resolve("Data!B1")
    assert result.status == "recovered"
    assert result.value == expected
