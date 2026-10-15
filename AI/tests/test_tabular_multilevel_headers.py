from openpyxl import Workbook

from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.tabular_validation import validated_table
from tests.test_tabular_report_recovery import _headers, _parsed


def test_vertically_merged_header_is_bound_to_its_actual_source_anchor():
    book = Workbook()
    sheet = book.active
    _headers(sheet, 8, 2, ["Entity", "Measurements"])
    sheet.merge_cells("B8:B9")
    sheet.merge_cells("C8:D8")
    _headers(sheet, 9, 3, ["Mass", "Length"])
    for row, values in [(10, ["A", 12, 30]), (11, ["B", 24, 60])]:
        for column, value in enumerate(values, 2):
            sheet.cell(row, column, value)
    summary, index = _parsed(book)
    table = build_query_tables(summary)[0]
    assert [column.header for column in table.columns] == ["Entity", "Mass", "Length"]
    assert table.columns[0].header_reference.endswith("!B8")
    assert [row.row_number for row in table.rows] == [10, 11]
    assert validated_table(table, index, {"B", "C", "D"})
