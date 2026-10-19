"""Source titles distinguish equal-valued histories belonging to different entities."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.agent.query.index import build_workbook_data_index
from app.agent.query.trend_sources import trend_candidates
from app.agent.query.trend_selection import distinct_trend_candidates, prepare_trend_candidate
from app.services.workbook_parser import parse_workbook


def test_distinct_company_captions_survive_identical_numeric_histories():
    book = Workbook()
    sheet = book.active
    for start, name in [(1, "Company North"), (8, "Company South")]:
        sheet.cell(start, 1, name).font = Font(bold=True)
        sheet.cell(start + 1, 1, "Date")
        sheet.cell(start + 1, 2, "Total Employees")
        for offset, (date, count) in enumerate([("2025-01-01", 10), ("2025-02-01", 20)], start + 2):
            sheet.cell(offset, 1, date)
            sheet.cell(offset, 2, count)
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    summary = parse_workbook("context.xlsx", content)
    index = build_workbook_data_index("context.xlsx", content)
    candidates = trend_candidates(summary, index, "전체 직원 수의 기간별 추이를 설명해줘")
    valid = [candidate for candidate in candidates if not candidate.error]
    assert len(valid) == 2
    names = [{cell.value for cell in candidate.context} for candidate in valid]
    assert any("Company North" in values for values in names)
    assert any("Company South" in values for values in names)
    prepared = [prepare_trend_candidate(candidate, index.rows) for candidate in valid]
    assert len(distinct_trend_candidates(prepared)) == 2
    assert {point.entity_cells[0].value for candidate in valid for point in candidate.observations} == {
        "Company North", "Company South"}
