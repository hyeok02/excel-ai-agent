from io import BytesIO
from xml.etree.ElementTree import fromstring, tostring
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from app.agent.query.formula_resolution import FormulaResolver
from app.agent.query.index import build_workbook_data_index


def source_book(formula="=SUM(Amounts)", *, headers=None, totals=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers or ["Region", "Amount", "Qty", "Calculated"])
    for values in [("East", 10, 2), ("West", 20, 3), ("East", 30, 4)]:
        sheet.append([*values, None])
    sheet["F1"] = formula
    if totals:
        sheet.append(["Total", 60, 9, None])
    sheet.add_table(Table(displayName="Sales", ref="A1:D5" if totals else "A1:D4",
                          totalsRowCount=1 if totals else 0))
    other = workbook.create_sheet("Other sheet")
    other["B2"] = 100
    hidden = workbook.create_sheet("Hidden")
    hidden["A1"] = 777
    hidden.sheet_state = "hidden"
    workbook.defined_names.add(DefinedName("Amounts", attr_text="'Data'!$B$2:$B$4"))
    workbook.defined_names.add(DefinedName("Regions", attr_text="'Data'!$A$2:$A$4"))
    workbook.defined_names.add(DefinedName("OneAmount", attr_text="'Data'!$B$2"))
    other.defined_names.add(DefinedName("Amounts", attr_text="'Other sheet'!$B$2"))
    return workbook


def content(workbook):
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def result(workbook, address="Data!F1", *, included=None, limits=None):
    index = build_workbook_data_index("any-file.xlsx", content(workbook), included)
    return FormulaResolver(index, limits).resolve(address)


def patch_xml(data, part, mutate):
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            body = source.read(entry)
            if entry.filename == part:
                root = fromstring(body)
                mutate(root)
                body = tostring(root)
            target.writestr(entry, body)
    return output.getvalue()
