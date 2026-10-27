"""Stored percentage series are neither invented ratios nor temporal buckets."""
import pytest
from openpyxl import Workbook

from app.agent.query.time_series_answers import time_series_answer
from app.agent.query.stored_ratio_phrases import stored_ratio_question
from app.agent.query.observed_series_answers import observed_series_average
from tests.support.large_workbook_fixtures import offline, parsed

LABEL = "Online sales as a percentage of total wholesale sales (%)"


def _fixture(fraction=False, rival=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Renamed observations"
    sheet.append(["Observation period", LABEL, "Revenue"])
    for month, value in zip(("Jan", "Feb", "Mar"), (10, 20, 30)):
        sheet.append([f"2023 {month}", value / 100 if fraction else value, 1000 * value])
    for row in range(2, 5):
        sheet.cell(row, 2).number_format = "0.0%" if fraction else "0.0"
    if rival:
        sheet["D1"] = "Online sales share of total wholesale sales (%)"
        for row in range(2, 5):
            sheet.cell(row, 4, 90)
    summary, index, _ = parsed(book)
    return summary, index


@pytest.mark.parametrize("fraction", [False, True])
def test_complete_stored_ratio_phrase_preserves_percent_scale_and_points(fraction):
    summary, index = _fixture(fraction)
    question = "Renamed observations 시트에서 2023년 1월부터 3월까지 온라인 도매판매 비중의 추이를 설명해줘"
    answer = time_series_answer(question, summary, index)
    assert answer is not None and answer.status == "answered"
    assert answer.verification.level == "request_checked"
    assert answer.question == question
    assert "3개 관측값" in answer.answer and "10% → 30%" in answer.answer
    assert "20%p 증가" in answer.answer
    assert "derived_trend" not in answer.selected_tools
    assert {item.reference for item in answer.evidence} >= {"B1", "B2", "B3", "B4"}


@pytest.mark.parametrize("fraction", [False, True])
def test_simple_mean_across_observations_is_one_number_not_monthly_buckets(fraction):
    summary, index = _fixture(fraction)
    question = f"Renamed observations 시트의 2023년 1월부터 3월까지 월별 {LABEL} 값의 단순 평균은 얼마야?"
    answer = time_series_answer(question, summary, index)
    assert answer is not None and answer.status == "answered"
    assert answer.verification.level == "request_checked"
    assert answer.selected_tools == ["observed_series_average"]
    assert "3개 단순 평균은 20%" in answer.answer
    assert "가중평균이 아닙니다" in answer.answer


def test_missing_months_do_not_certify_a_full_year_mean():
    summary, index = _fixture()
    answer = time_series_answer(f"Renamed observations 시트의 2023년 월별 {LABEL} 값의 단순 평균은 얼마야?", summary, index)
    assert answer is not None and answer.status == "limited"
    assert answer.verification.level == "source_checked"
    assert any("일부 관측 시점" in reason for reason in answer.limitations)


@pytest.mark.parametrize("phrase", ["온라인 소매판매 비중", "온라인 도매판매액", "도매판매 비중", "온라인 도매판매 비중상한"])
def test_partial_wrong_metric_or_unknown_qualifier_is_not_bound(phrase):
    summary, _ = _fixture()
    question = phrase + " 추이"
    assert stored_ratio_question(question, summary) == question


def test_semantically_different_source_headers_with_same_atoms_remain_ambiguous():
    summary, _ = _fixture(rival=True)
    question = "온라인 도매판매 비중 추이"
    assert stored_ratio_question(question, summary) == question


def test_parent_whole_and_part_roles_cannot_be_swapped_by_matching_word_sets():
    from app.agent.query import stored_ratio_phrases
    from app.agent.query.tabular_catalog import build_query_tables
    summary, _ = _fixture()
    tables = build_query_tables(summary)
    column = next(column for column in tables[0].columns if column.id == "B")
    assert stored_ratio_phrases._label_atoms(column.header) != stored_ratio_phrases._label_atoms(
        "Wholesale sales as a percentage of total online sales (%)")


@pytest.mark.parametrize("qualifier", ["10% 이상인", "지역별", "가중", "전년 대비"])
def test_unverified_filters_or_extra_computations_cannot_be_erased_by_mean(qualifier):
    summary, index = _fixture()
    question = f"2023년 1월부터 3월까지 {LABEL} {qualifier} 값의 단순 평균"
    answer = observed_series_average(question, summary, index)
    assert answer is None or answer.verification.level != "request_checked"
