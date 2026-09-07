"""Edge cases for source-only ranked table narratives."""

from io import BytesIO

import pytest
from openpyxl import Workbook

from app.services.insights.context import build_workbook_context
from app.services.insights.ranked_narratives import ranked_report
from app.services.insights.table_schema import share_value
from app.services.region_detector import CellRegion
from app.services.workbook_details.regions import summarize_regions
from app.services.workbook_parser import parse_workbook


def _cell(address, value, number_format="General"):
    return {"cell": address, "value": value, "number_format": number_format}


def _ranked_sheet(name, holder, percent):
    rows = [
        [_cell("A1", "Rank"), _cell("B1", "Name"),
         _cell("C1", "Shares Held"), _cell("D1", "Share (%)")],
        [_cell("A2", 1), _cell("B2", holder),
         _cell("C2", 1_000), _cell("D2", percent)],
    ]
    return {
        "name": name,
        "business_facts": {
            "table_regions": [{"title": "Top Holders", "rows": rows}],
        },
    }


def test_formula_backed_region_title_uses_cached_text_before_rank_value():
    workbook, values = Workbook(), Workbook()
    sheet, value_sheet = workbook.active, values.active
    sheet["A1"] = '="Top Holders"'
    value_sheet["A1"] = "Top Holders"
    sheet.append([1, "Primary Fund", 42])
    try:
        summary = summarize_regions(
            sheet, [CellRegion("A1", "C2", 4)], value_sheet,
        )[0]
        assert summary.title == "Top Holders"
        assert summary.title != "1"
    finally:
        workbook.close()
        values.close()


def test_percentage_number_format_normalizes_decimal_fraction():
    cell = {"value": 0.6581, "number_format": "0.00%"}
    assert share_value(cell) == pytest.approx(65.81)


def test_primary_ranked_sheet_precedes_later_peer_with_larger_percentage():
    context = {
        "sheets": [
            _ranked_sheet("Peer Comparison", "Peer Fund", 95),
            _ranked_sheet("Ownership", "Primary Fund", 42),
        ],
    }

    insights, _ = ranked_report(context)

    assert "Primary Fund" in insights[0].fact
    assert all("Peer Fund" not in item.fact for item in insights)


def test_peer_name_with_underscore_is_excluded_without_outline_titles():
    peer = _ranked_sheet("Peer_Ownership", "Peer Fund", 95)
    peer["content_outline"] = {"region_titles": []}
    context = {"sheets": [_ranked_sheet("Ownership", "Primary Fund", 42), peer]}

    insights, _ = ranked_report(context)

    assert insights and all("Peer Fund" not in item.fact for item in insights)


def test_comparison_table_is_used_when_it_is_the_only_data_table():
    context = {
        "sheets": [
            {"name": "Instructions", "business_facts": {}},
            _ranked_sheet("Peer_Ownership", "Only Fund", 42),
        ],
    }

    insights, _ = ranked_report(context)

    assert insights and "Only Fund" in insights[0].fact


def test_ranked_table_title_is_used_only_with_its_source_cell():
    sheet = _ranked_sheet("Ownership", "Primary Fund", 42)
    region = sheet["business_facts"]["table_regions"][0]
    region["title_cell"] = _cell("A5", "Top Holders")

    insights, _ = ranked_report({"sheets": [sheet]})

    assert "Top Holders에서" in insights[0].fact
    assert "'Ownership'!A5" in insights[0].evidence


def test_parser_separated_ranked_header_and_body_are_rejoined():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Rank", "Name", "Shares Held", "Share (%)"])
    sheet.append([1, "Alpha Fund", 100, 42])
    sheet.append([2, "Beta Fund", 80, 33])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()

    summary = parse_workbook("plain-ranked.xlsx", stream.getvalue())
    insights, _ = ranked_report(build_workbook_context(summary))

    assert insights and "Alpha Fund" in insights[0].fact
    assert "42%" in insights[0].fact
    assert "Rank에서" not in insights[0].fact


def test_ranked_header_does_not_cross_a_new_section_title():
    header = {"rows": [[
        _cell("A1", "Rank"), _cell("B1", "Name"),
        _cell("C1", "Shares Held"), _cell("D1", "Share (%)"),
    ]]}
    other_section = {"rows": [
        [_cell("A2", "Peer Holdings")],
        [_cell("A3", 1), _cell("B3", "Wrong Fund"),
         _cell("C3", 100), _cell("D3", 42)],
    ]}
    context = {"sheets": [{"name": "Ownership", "business_facts": {
        "table_regions": [header, other_section],
    }}]}

    insights, _ = ranked_report(context)

    assert all("Wrong Fund" not in item.fact for item in insights)
