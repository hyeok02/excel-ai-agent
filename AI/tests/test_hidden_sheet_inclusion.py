"""숨김 시트를 분석에 넣는 선택지 검증.

숨김 시트는 기본적으로 빼지만, 계산 과정을 숨겨 두는 실무 파일이 있어
사용자가 켜면 보이는 시트와 똑같이 다룬다.
"""
from openpyxl import Workbook

from app.services.analysis_inclusion import AnalysisDecision
from app.services.workbook_parser import parse_workbook
from app.services.worksheet_filter import evaluate_worksheet_inclusion


def workbook_bytes() -> bytes:
    from io import BytesIO

    book = Workbook()
    visible = book.active
    visible.title = "요약"
    visible["A1"] = "부서"
    visible["A2"] = "서비스"
    visible["B1"] = "인원"
    visible["B2"] = 1081
    hidden = book.create_sheet("계산과정")
    hidden["A1"] = "계수"
    hidden["A2"] = 0.9
    hidden.sheet_state = "hidden"
    cache = book.create_sheet("___snlofficequeries")
    cache["A1"] = "cache"
    cache.sheet_state = "hidden"
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def sheet(name: str, state: str = "visible"):
    book = Workbook()
    worksheet = book.active
    worksheet.title = name
    worksheet.sheet_state = state
    return worksheet


def test_hidden_sheets_stay_out_by_default():
    found = evaluate_worksheet_inclusion(sheet("계산과정", "hidden"))

    assert found.decision is AnalysisDecision.EXCLUDE
    assert found.reason_code == "hidden_worksheet"


def test_hidden_sheets_are_analysed_when_the_user_turns_it_on():
    found = evaluate_worksheet_inclusion(sheet("계산과정", "hidden"), include_hidden=True)

    assert found.decision is AnalysisDecision.INCLUDE


def test_add_in_cache_sheets_stay_out_even_when_turned_on():
    # 사용자가 만든 내용이 아니라 애드인이 남긴 캐시라 켜도 분석하지 않는다.
    found = evaluate_worksheet_inclusion(
        sheet("___snlofficequeries", "hidden"), include_hidden=True
    )

    assert found.decision is AnalysisDecision.EXCLUDE
    assert found.reason_code == "addin_cache_worksheet"


def test_very_hidden_sheets_follow_the_same_switch():
    assert evaluate_worksheet_inclusion(
        sheet("내부", "veryHidden")
    ).decision is AnalysisDecision.EXCLUDE
    assert evaluate_worksheet_inclusion(
        sheet("내부", "veryHidden"), include_hidden=True
    ).decision is AnalysisDecision.INCLUDE


def test_parsing_counts_the_hidden_sheet_only_when_turned_on():
    content = workbook_bytes()

    default = parse_workbook("book.xlsx", content)
    opened = parse_workbook("book.xlsx", content, include_hidden=True)

    assert [item.name for item in default.sheets] == ["요약"]
    assert [item.name for item in opened.sheets] == ["요약", "계산과정"]
    assert default.excluded_sheet_count == 2
    assert opened.excluded_sheet_count == 1
