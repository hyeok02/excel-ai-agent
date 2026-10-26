"""Aliases preserve source meaning across renamed, merged and plain reports."""
import asyncio

from openpyxl import Workbook
from openpyxl.styles import Font
import pytest

from app.agent import create_default_tool_registry
from app.agent.query.service import WorkbookQuestionService
from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.metric_semantics import column_metric_conflict
from tests.support.large_workbook_fixtures import NoGeneratedAnswer, parsed


def source(*, parent="Roles", merged=False, financial=False, category=True):
    book = Workbook()
    sheet = book.active
    sheet.title = "Arbitrary source"
    if category:
        sheet.append([None, parent, None])
        sheet["B1"].font = Font(bold=True)
        if merged:
            sheet.merge_cells("B1:C1")
        sheet.append(["Date", "Support", "Sales"])
        for cell in sheet[2]:
            cell.font = Font(bold=True)
        sheet.append(["2024-01-01", 8, 2])
        sheet.append(["2025-01-01", 9, 4])
        for row in sheet.iter_rows(min_row=3):
            for cell in row[1:]:
                cell.number_format = '0"명"'
    if financial or not category:
        money = book.create_sheet("Amounts") if category else sheet
        money.append(["Date", "Sales"])
        money.append(["2024-01-01", 100])
        money.append(["2025-01-01", 150])
        for row in money.iter_rows(min_row=2):
            row[1].number_format = '"USD "#,##0'
    return parsed(book, filename="not-a-known-file.xlsx")[:2]


def answer(question, **kwargs):
    summary, index = source(**kwargs)
    service = WorkbookQuestionService(NoGeneratedAnswer(), create_default_tool_registry())
    return asyncio.run(service.answer(question, summary, index))


@pytest.mark.parametrize("question", ["매출 추이 알려줘", "Revenue 추이 알려줘", "매출 합계 알려줘"])
@pytest.mark.parametrize("merged", [False, True])
def test_category_alias_never_answers_a_financial_measure(question, merged):
    result = answer(question, merged=merged)
    assert result.status == "insufficient_evidence"
    assert result.verification.level == "unverified"
    assert result.selected_tools == ["check_metric_semantics"]
    assert "2 → 4" not in result.answer


@pytest.mark.parametrize("parent", ["Roles", "Department", "직무", "부서"])
def test_parent_binding_is_source_based_not_sheet_or_file_specific(parent):
    summary, _ = source(parent=parent)
    sales = next(c for t in build_query_tables(summary) for c in t.columns if c.header == "Sales")
    assert any(text == parent for _, text in sales.context)
    assert column_metric_conflict(sales, "매출 추이")


@pytest.mark.parametrize("question", ["Sales 추이 알려줘", "Sales 직무 인원 추이 알려줘"])
def test_explicit_category_queries_remain_answerable(question):
    result = answer(question)
    assert result.status == "answered", result.model_dump()
    assert "2명 → 4명" in result.answer


@pytest.mark.parametrize("financial", [False, True])
def test_financial_sales_alias_still_works_and_ignores_role_table(financial):
    result = answer("매출 추이 알려줘", category=financial, financial=financial)
    assert result.status == "answered", result.model_dump()
    assert "100" in result.answer and "150" in result.answer
    assert "2명" not in result.answer
