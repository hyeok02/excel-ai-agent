"""Rotating a table or adding unrelated merges cannot change metric meaning."""
import asyncio

from openpyxl import Workbook
from openpyxl.styles import Font
import pytest

from app.agent import create_default_tool_registry
from app.agent.query.metric_observations import horizontal_observations
from app.agent.query.service import WorkbookQuestionService
from tests.support.large_workbook_fixtures import NoGeneratedAnswer, parsed


def source(*, horizontal=True, money=False, roles=True):
    book = Workbook()
    sheet = book.active
    sheet.title = "Another arbitrary name"
    if roles and horizontal:
        sheet.append(["Roles"])
        sheet.append(["Role", "2024-01-01", "2025-01-01"])
        sheet.append(["Sales", 2, 4])
        sheet.append(["Support", 8, 9])
    elif roles:
        sheet.append([None, "Roles", None, None, None, "Notes"])
        sheet.merge_cells("F1:G1")
        sheet.append(["Date", "Support", "Sales"])
        sheet.append(["2024-01-01", 8, 2])
        sheet.append(["2025-01-01", 9, 4])
    if roles:
        for row in sheet.iter_rows(min_row=1, max_row=2):
            for cell in row:
                cell.font = Font(bold=True)
        for row in sheet.iter_rows(min_row=3, min_col=2, max_col=3):
            for cell in row:
                cell.number_format = '0"명"'
    if money:
        financial = book.create_sheet("Financial") if roles else sheet
        financial.append(["Measure", "2024-01-01", "2025-01-01"])
        financial.append(["Sales", 100, 150])
        for cell in financial[1]:
            cell.font = Font(bold=True)
        for cell in financial[2][1:]:
            cell.number_format = '"USD "#,##0'
    return parsed(book, filename="arbitrary.xlsx")[:2]


def answer(question, **kwargs):
    summary, index = source(**kwargs)
    service = WorkbookQuestionService(NoGeneratedAnswer(), create_default_tool_registry())
    return asyncio.run(service.answer(question, summary, index))


@pytest.mark.parametrize("horizontal", [False, True])
@pytest.mark.parametrize("question", ["매출 추이 알려줘", "Revenue 추이 알려줘",
                                      "Sales 부서의 Revenue 추이 알려줘", "Sales 직무의 매출 추이를 알려줘"])
def test_role_alias_never_becomes_a_verified_financial_answer(horizontal, question):
    result = answer(question, horizontal=horizontal)
    assert result.status == "insufficient_evidence"
    assert result.verification.level == "unverified"
    assert result.selected_tools == ["check_metric_semantics"]


@pytest.mark.parametrize("question", ["Sales 추이 알려줘", "Sales 직무 인원 추이 알려줘"])
def test_explicit_horizontal_category_still_works(question):
    result = answer(question)
    assert result.status == "answered", result.model_dump()
    assert "2명 → 4명" in result.answer


@pytest.mark.parametrize("roles", [False, True])
def test_horizontal_financial_sales_remains_available_alongside_roles(roles):
    result = answer("매출 추이 알려줘", roles=roles, money=True)
    assert result.status == "answered", result.model_dump()
    assert "100" in result.answer and "150" in result.answer
    assert "2명" not in result.answer


def test_observation_and_clarification_paths_exclude_incompatible_role_values():
    _, index = source(money=True)
    points = horizontal_observations(index.rows, "매출 추이 알려줘")
    assert {point.value.value for point in points} == {100, 150}


@pytest.mark.parametrize("header", ["Category", "Categories", "분류", "항목"])
@pytest.mark.parametrize("question,expected", [("매출 추이 알려줘", "100"), ("영업이익 추이 알려줘", "20")])
def test_generic_category_header_does_not_disqualify_financial_metrics(header, question, expected):
    book = Workbook()
    sheet = book.active
    sheet.append([header, "2024-01-01", "2025-01-01"])
    sheet.append(["Revenue", 100, 150])
    sheet.append(["Operating Income", 20, 30])
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = '"USD "#,##0'
    summary, index = parsed(book, filename="finance.xlsx")[:2]
    service = WorkbookQuestionService(NoGeneratedAnswer(), create_default_tool_registry())
    result = asyncio.run(service.answer(question, summary, index))
    assert result.status == "answered", result.model_dump()
    assert expected in result.answer
