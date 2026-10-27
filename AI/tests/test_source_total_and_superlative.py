"""Whole-population and top-one requests use source fields, never fixture names."""
from dataclasses import replace

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from app.agent.query.source_aggregate_answers import source_aggregate_answer
from app.agent.query.tabular_catalog import build_query_tables
from app.agent.query.tabular_execution import execute_tabular_query
from app.agent.query.tabular_query_models import TabularFilter, TabularQueryPlan
from tests.support.large_workbook_fixtures import parsed


def fixture(start=2, column=1, reverse=False, total_delta=0):
    book = Workbook()
    sheet = book.active
    sheet.title = "Fresh statistics"
    sheet.cell(start, column, "지역")
    sheet.merge_cells(start_row=start, start_column=column, end_row=start + 1, end_column=column)
    groups = ["시설수", "이용자수"]
    if reverse:
        groups.reverse()
    for offset, label in enumerate(groups):
        first = column + 1 + offset * 3
        sheet.cell(start, first, label)
        sheet.merge_cells(start_row=start, start_column=first, end_row=start, end_column=first + 2)
        for delta, child in enumerate(["계", "공립", "사립"]):
            sheet.cell(start + 1, first + delta, child)
    rows = [("총합계", {"시설수": [30, 10, 20], "이용자수": [1020 + total_delta, 420, 600]}),
            ("동부", {"시설수": [12, 4, 8], "이용자수": [320, 120, 200]}),
            ("서부", {"시설수": [18, 6, 12], "이용자수": [700, 300, 400]})]
    for offset, (label, values) in enumerate(rows, 2):
        for delta, value in enumerate([label, *[value for group in groups for value in values[group]]]):
            sheet.cell(start + offset, column + delta, value)
    for row in sheet.iter_rows(min_row=start, max_row=start + 1):
        for cell in row:
            cell.font = Font(bold=True)
    summary, index, _ = parsed(book)
    return summary, index, build_query_tables(summary)[0]


def named(table, text):
    return next(column.id for column in table.columns if column.display_header == text)


@pytest.mark.parametrize("start,column,reverse", [(2, 1, False), (14, 7, True)])
def test_superlative_uses_parent_total_and_excludes_named_total(start, column, reverse):
    summary, index, table = fixture(start, column, reverse)
    identity = table.columns[0].id
    metric = named(table, "시설수 / 계")
    plan = TabularQueryPlan(operation="row_list", table_id=table.id,
        column_ids=[identity, metric], order_by=metric, limit=1,
        filters=[TabularFilter(column_id=identity, operator="not_equals", value="총합계")])
    result = execute_tabular_query(plan, summary, index, "총합계를 제외하고 시설수가 가장 많은 지역은 어디야?")
    assert result is not None and result.status == "answered"
    assert "지역: 서부" in result.answer and "시설수 / 계" in result.answer and "계: 18" in result.answer
    assert result.verification.level == "request_checked"
    for wrong in [named(table, "이용자수 / 계"), named(table, "시설수 / 사립")]:
        forged = plan.model_copy(update={"column_ids": [identity, wrong], "order_by": wrong})
        assert execute_tabular_query(forged, summary, index,
            "총합계를 제외하고 시설수가 가장 많은 지역은 어디야?") is None


@pytest.mark.parametrize("question", ["전국 이용자수는 총 몇 명이야?", "전체 이용자수는 얼마야?"])
@pytest.mark.parametrize("start,column,reverse", [(2, 1, False), (14, 7, True)])
def test_whole_scope_can_select_unique_reconciled_total(question, start, column, reverse):
    summary, index, table = fixture(start, column, reverse)
    result = source_aggregate_answer(question, summary, index)
    assert result is not None and "이용자수 / 계: 1,020" in result.answer
    assert result.verification.level == "request_checked"
    # The same independent scope check accepts a model selecting that exact cell.
    metric = named(table, "이용자수 / 계")
    plan = TabularQueryPlan(operation="lookup", table_id=table.id, column_ids=[metric],
        filters=[TabularFilter(column_id=table.columns[0].id, operator="equals", value="총합계")])
    assert execute_tabular_query(plan, summary, index, question) is not None
    wrong = plan.model_copy(update={"column_ids": [named(table, "시설수 / 계")]})
    assert execute_tabular_query(wrong, summary, index, question) is None


@pytest.mark.parametrize("question", [
    "전국 이용자수는 총 몇 명이야? 서울만", "전국 이용자수는 총 몇 명이야? 2025년만",
    "전국 이용자수는 총 몇 명이야? 2배 해줘", "전국 이용자수와 시설수는 총 몇 명이야?",
])
def test_total_alias_never_discards_unknown_scope(question):
    summary, index, _ = fixture()
    assert source_aggregate_answer(question, summary, index) is None


def test_non_reconciling_or_incomplete_totals_cannot_authorize_alias():
    summary, index, _ = fixture(total_delta=1)
    assert source_aggregate_answer("전국 이용자수는 총 몇 명이야?", summary, index) is None
    summary, index, table = fixture()
    from app.agent.query.tabular_total_scope import total_lookup_plan
    assert total_lookup_plan("전국 이용자수는 총 몇 명이야?", replace(table, complete=False)) is None
    assert total_lookup_plan("전국 이용자수는 총 몇 명이야?", replace(table, excluded_rows=frozenset({9}))) is None


def test_total_and_details_cannot_be_added_twice():
    summary, index, table = fixture()
    metric = named(table, "이용자수 / 계")
    plan = TabularQueryPlan(operation="numeric_summary", table_id=table.id,
                            column_ids=[metric], statistic="sum")
    result = execute_tabular_query(plan, summary, index, "이용자수 합계 알려줘")
    assert result is None or result.status == "insufficient_evidence"


@pytest.mark.parametrize("metric", ["Highest Score", "Lowest Cost", "가장 높은 지점"])
def test_superlative_inside_source_header_is_not_an_ordering_request(metric):
    from tests.tabular_query_fixtures import table_fixture
    summary, index = table_fixture([(2, 2, [["Item", metric], ["West", 12], ["East", 9]])])
    table = build_query_tables(summary)[0]
    plan = TabularQueryPlan(operation="lookup", table_id=table.id, column_ids=["C"],
        filters=[TabularFilter(column_id="B", operator="equals", value="West")])
    result = execute_tabular_query(plan, summary, index, f"West의 {metric} 값은?")
    assert result is not None and result.verification.level == "request_checked"


def test_total_alias_is_not_a_license_for_join_filter_mutation():
    from types import SimpleNamespace
    from app.agent.query.tabular_total_scope import whole_total_filter_supported
    _, _, table = fixture()
    condition = TabularFilter(column_id="A", operator="equals", value="동부")
    joined = SimpleNamespace(operation="numeric_summary", statistic="sum", filters=[condition])
    assert not whole_total_filter_supported("전체 이용자수 합계", joined, table, condition)
