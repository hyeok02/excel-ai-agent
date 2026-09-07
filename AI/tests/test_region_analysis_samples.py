from datetime import date

from openpyxl import Workbook

from app.services.region_detector import CellRegion
from app.services.workbook_details.analysis_samples import collect_analysis_rows
from app.services.workbook_details.regions import summarize_regions


def test_small_meal_region_includes_lower_numeric_rows_and_all_days() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["H6"] = "월요일 식단"
    sheet["C21"] = "영양소"
    sheet["C22"] = "에너지(kcal)"
    sheet["H22"] = 678.74
    sheet["N22"] = 698.05
    sheet["C31"] = "철분(mg)"
    sheet["N31"] = 10.71
    try:
        region = summarize_regions(sheet, [CellRegion("C6", "O31", 7)])[0]
        assert len(region.preview_rows) == 8
        assert all(len(row) == 8 for row in region.preview_rows)
        assert region.preview_rows[-1][-1].address == "J13"
        assert region.is_truncated
        assert len(region.analysis_rows) == 26
        assert all(len(row) == 13 for row in region.analysis_rows)
        cells = {cell["address"]: cell for row in region.analysis_rows for cell in row}
        assert cells["H22"]["value"] == 678.74
        assert cells["N22"]["value"] == 698.05
        assert cells["N31"]["value"] == 10.71
        assert set(cells["N31"]) == {
            "address", "value", "formula", "cached_value", "number_format",
        }
    finally:
        workbook.close()


def test_large_regions_reuse_the_existing_preview_without_extra_snapshots() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "대규모 영역"
    sheet["Z100"] = 10
    try:
        region = summarize_regions(sheet, [CellRegion("A1", "Z100", 2)])[0]
        assert region.analysis_rows == []
        assert len(region.preview_rows) == 8
        assert all(len(row) == 8 for row in region.preview_rows)
    finally:
        workbook.close()


def test_large_rectangle_is_rejected_before_any_cell_access() -> None:
    class NoCellAccess:
        def cell(self, **_):
            raise AssertionError("oversized sample must not access worksheet cells")

    assert collect_analysis_rows(NoCellAccess(), None, (1, 1, 16384, 1048576), 2048) == []


def test_total_compact_samples_are_bounded_per_sheet() -> None:
    workbook = Workbook()
    sheet = workbook.active
    regions = []
    for offset in range(5):
        start = 1 + offset * 65
        sheet.cell(start, 1, "기록")
        sheet.cell(start + 63, 8, offset)
        regions.append(CellRegion(f"A{start}", f"H{start + 63}", 2))
    try:
        summaries = summarize_regions(sheet, regions)
        counts = [sum(map(len, region.analysis_rows)) for region in summaries]
        assert counts == [512, 512, 512, 512, 0]
        assert sum(counts) == 2048
        assert all(len(region.preview_rows) == 8 for region in summaries)
    finally:
        workbook.close()


def test_compact_samples_preserve_formula_cache_dates_and_number_format() -> None:
    workbook, values = Workbook(), Workbook()
    sheet, value_sheet = workbook.active, values.active
    sheet["A1"] = date(2026, 9, 1)
    sheet["B1"] = "=SUM(C1:D1)"
    sheet["B1"].number_format = "0.00"
    value_sheet["B1"] = 12.5
    try:
        rows = collect_analysis_rows(sheet, value_sheet, (1, 1, 2, 1), 2048)
        assert rows[0][0]["value"] == "2026-09-01"
        assert rows[0][1] == {
            "address": "B1", "value": None, "formula": "=SUM(C1:D1)",
            "cached_value": 12.5, "number_format": "0.00",
        }
        assert sheet["B1"].value == "=SUM(C1:D1)"
    finally:
        workbook.close()
        values.close()
