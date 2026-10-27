"""Repeated leaves do not allow a source-valid value from the wrong measure."""
from dataclasses import replace

import pytest
from openpyxl import Workbook

from app.agent.query.intent_models import QuestionIntent
from app.agent.query.tabular_catalog import build_query_tables, query_catalog
from app.agent.query.tabular_execution import execute_tabular_query
from app.agent.query.tabular_query_models import TabularFilter, TabularQueryPlan
from app.agent.query.tabular_validation import validated_table
from tests.test_tabular_report_recovery import _headers
from tests.support.large_workbook_fixtures import offline, parsed


def _fixture(start_row=2, start_column=1, reversed_groups=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Renamed hierarchical report"
    groups = ["시설수", "이용자수"]
    if reversed_groups:
        groups.reverse()
    _headers(sheet, start_row, start_column, ["지역"])
    sheet.merge_cells(start_row=start_row, start_column=start_column,
                      end_row=start_row + 1, end_column=start_column)
    _headers(sheet, start_row + 1, start_column + 1, ["계", "공립", "사립"] * 2)
    for offset, name in enumerate(groups):
        first = start_column + 1 + 3 * offset
        _headers(sheet, start_row, first, [name])
        sheet.merge_cells(start_row=start_row, start_column=first,
                          end_row=start_row, end_column=first + 2)
    for row_offset, (name, values) in enumerate([
        ("동부", {"시설수": [12, 4, 8], "이용자수": [320, 120, 200]}),
        ("서부", {"시설수": [18, 6, 12], "이용자수": [700, 300, 400]}),
    ], 2):
        for offset, value in enumerate([name] + [value for group in groups for value in values[group]]):
            sheet.cell(start_row + row_offset, start_column + offset, value)
    summary, index, _ = parsed(book)
    table = build_query_tables(summary)[0]
    return summary, index, table


def _answer(summary, index, table, question, selected):
    plan = TabularQueryPlan(operation="lookup", table_id=table.id, column_ids=[selected], filters=[
        TabularFilter(column_id=table.columns[0].id, operator="equals", value="동부"),
    ])
    return execute_tabular_query(plan, summary, index, question,
                                 QuestionIntent(False, "cell_lookup", (), True, plan))


@pytest.mark.parametrize("start_row,start_column,reversed_groups", [(2, 1, False), (13, 7, True)])
def test_requested_parent_and_child_bind_after_renaming_relocation_and_swapping(start_row, start_column, reversed_groups):
    summary, index, table = _fixture(start_row, start_column, reversed_groups)
    correct = next(column for column in table.columns if column.display_header == "시설수 / 계")
    wrong_parent = next(column for column in table.columns if column.display_header == "이용자수 / 계")
    wrong_child = next(column for column in table.columns if column.display_header == "시설수 / 사립")
    public = next(column for column in table.columns if column.display_header == "시설수 / 공립")

    result = _answer(summary, index, table, "동부의 시설수는 모두 몇 개야?", correct.id)
    assert result is not None and "시설수 / 계: 12" in result.answer
    assert result.verification.level == "request_checked"
    assert _answer(summary, index, table, "동부의 시설수는 모두 몇 개야?", wrong_parent.id) is None
    assert _answer(summary, index, table, "동부의 공립 시설수는 몇 개야?", wrong_child.id) is None
    assert _answer(summary, index, table, "동부의 공립시설수는 몇 개야?", wrong_parent.id) is None
    result = _answer(summary, index, table, "동부의 공립 시설수는 몇 개야?", public.id)
    assert result is not None and "시설수 / 공립: 4" in result.answer
    assert result.verification.level == "request_checked"


def test_qualified_header_is_in_catalog_and_parent_is_read_back_not_trusted():
    summary, index, table = _fixture()
    column = next(column for column in table.columns if column.display_header == "시설수 / 계")
    catalog = query_catalog(summary, "시설수")
    item = next(entry for entry in catalog[0]["columns"] if entry["column_id"] == column.id)
    assert item["qualified_header"] == "시설수 / 계"
    assert item["header_path"] == [("Renamed hierarchical report!B2", "시설수")]
    assert validated_table(table, index, {column.id})
    tampered = replace(column, header_path=((column.header_path[0][0], "other measure"),))
    forged = replace(table, columns=tuple(tampered if item.id == column.id else item for item in table.columns))
    assert not validated_table(forged, index, {column.id})


def test_source_stream_rechecks_merged_anchors_and_keeps_parent_scopes():
    summary, index, table = _fixture()
    index = replace(index, rows=tuple(row for row in index.rows if row.row_number <= 3),
                    truncated=True, truncated_sheet_names=frozenset({table.sheet_name}))
    correct = next(column for column in table.columns if column.display_header == "시설수 / 계")
    wrong = next(column for column in table.columns if column.display_header == "이용자수 / 계")
    result = _answer(summary, index, table, "동부의 시설수는 모두 몇 개야?", correct.id)
    assert result is not None and "시설수 / 계: 12" in result.answer
    assert "full_source_scan" in result.verification.checks
    assert _answer(summary, index, table, "동부의 시설수는 모두 몇 개야?", wrong.id) is None


def test_qualified_header_can_bind_but_duplicate_leaf_alone_cannot():
    summary, index, table = _fixture()
    assert _answer(summary, index, table, "동부의 시설수는 모두 몇 개야?", "계") is None
    result = _answer(summary, index, table, "동부의 시설수는 모두 몇 개야?", "시설수 / 계")
    assert result is not None and "시설수 / 계: 12" in result.answer
