from io import BytesIO

import pytest
from openpyxl import Workbook

from tests.support.writeback_target_fixtures import proposal


def source(rows=None, duplicate_sheet=False):
    book = Workbook()
    book.active.title = "Orders"
    rows = rows or [
        ["Region", "Product", "Revenue", "Quantity"],
        ["Seoul", "Notebook", 10, 2],
        ["Seoul", "Monitor", 20, 3],
        ["Busan", "Notebook", 30, 4],
    ]
    for row in rows:
        book.active.append(row)
    if duplicate_sheet:
        book.copy_worksheet(book.active).title = "Other"
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


@pytest.mark.parametrize("instruction,reference", [
    ("Orders Seoul Notebook Revenue 값을 99로 수정", "C3"),
    ("Orders Seoul Notebook Revenue 값을 99로 수정", "C4"),
    ("Orders Notebook Revenue 값을 99로 수정", "C2"),
    ("Orders Seoul Notebook 값을 99로 수정", "C2"),
    ("Orders UnknownMetric 값을 99로 수정", "C2"),
    ("Orders NotebookPro Revenue 값을 99로 수정", "C2"),
])
def test_rejects_unproven_or_ambiguous_natural_targets(instruction, reference):
    body = proposal(instruction, [("Orders", reference, 99)], source())
    assert body["status"] == "blocked", body
    assert not body["changes"]


@pytest.mark.parametrize("reference", ["C2", "C3", "C4"])
def test_all_rows_does_not_remove_other_filters(reference):
    body = proposal("Orders Seoul의 모든 Revenue 값을 99로 수정",
                    [("Orders", reference, 99)], source())
    assert (body["status"] == "ready") == (reference in {"C2", "C3"}), body


@pytest.mark.parametrize("instruction,references", [
    ("Orders Seoul Notebook Revenue 값을 99로 수정", ["C2"]),
    ("Orders Seoul Notebook와 Monitor Revenue 값을 99로 수정", ["C2", "C3"]),
    ("Orders 모든 Revenue 값을 99로 수정", ["C2", "C3", "C4"]),
    ("Orders C2:C4를 99로 수정", ["C2", "C3", "C4"]),
])
def test_preserves_intersection_enumeration_and_explicit_ranges(instruction, references):
    body = proposal(instruction, [("Orders", ref, 99) for ref in references], source())
    assert body["status"] == "ready", body
    assert len(body["changes"]) == len(references)


def test_duplicate_named_items_across_sheets_require_sheet_scope():
    body = proposal("Seoul Notebook Revenue 값을 99로 수정", [("Orders", "C2", 99)],
                    source(duplicate_sheet=True))
    assert body["status"] == "blocked", body
    named = proposal("Orders Seoul Notebook Revenue 값을 99로 수정",
                     [("Orders", "C2", 99)], source(duplicate_sheet=True))
    assert named["status"] == "ready", named


def test_named_row_with_one_value_is_not_unnecessarily_blocked():
    body = proposal("Orders Notebook 값을 99로 수정", [("Orders", "B2", 99)],
                    source([["Product", "Amount"], ["Notebook", 10], ["Monitor", 20]]))
    assert body["status"] == "ready", body


@pytest.mark.parametrize("reference,value,ready", [
    ("C2", 99, True), ("C4", 55, True), ("C2", 55, False), ("C4", 99, False),
])
def test_each_natural_destination_keeps_its_own_value(reference, value, ready):
    body = proposal("Orders Seoul Notebook Revenue를 99로, Busan Notebook Revenue를 55로 수정",
                    [("Orders", reference, value)], source())
    assert (body["status"] == "ready") == ready, body


def test_source_label_is_not_a_replacement_value():
    body = proposal("Orders Seoul Notebook Revenue를 99로 수정",
                    [("Orders", "C2", "Notebook")], source())
    assert body["status"] == "blocked", body


def test_shared_headers_do_not_make_a_unique_row_ambiguous_across_sheets():
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(source()))
    other = book.create_sheet("Other")
    other.append(["Region", "Product", "Revenue", "Quantity"])
    other.append(["Jeju", "Keyboard", 30, 4])
    output = BytesIO()
    book.save(output)
    book.close()
    body = proposal("Seoul Notebook Revenue 값을 99로 수정", [("Orders", "C2", 99)], output.getvalue())
    assert body["status"] == "ready", body


def test_a_prior_tables_header_does_not_authorize_another_metric():
    content = source([["Product", "Revenue"], ["Notebook", 10], [],
                      ["Product", "Cost"], ["Monitor", 20]])
    body = proposal("Orders Monitor Revenue 값을 99로 수정", [("Orders", "B5", 99)], content)
    assert body["status"] == "blocked", body
    valid = proposal("Orders Monitor Cost 값을 99로 수정", [("Orders", "B5", 99)], content)
    assert valid["status"] == "ready", valid


def test_a_single_existing_row_is_not_a_substitute_for_a_missing_requested_row():
    content = source([["Product", "Amount"], ["Notebook", 10]])
    missing = proposal("Orders Monitor Amount 값을 99로 수정", [("Orders", "B2", 99)], content)
    assert missing["status"] == "blocked", missing
    unique = proposal("Orders Amount 값을 99로 수정", [("Orders", "B2", 99)], content)
    assert unique["status"] == "ready", unique
